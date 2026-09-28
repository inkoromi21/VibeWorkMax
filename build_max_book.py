from __future__ import annotations

from pathlib import Path
from textwrap import wrap
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/inkoromi21/Documents/VibeWorkMax')
OUT = ROOT / 'MAX_образовательная_платформа_книга_разработчика.docx'
ASSETS = ROOT / 'doc_assets'
ASSETS.mkdir(exist_ok=True)

BLUE = '163A5F'
MID_BLUE = 'DCEAF6'
PALE_BLUE = 'F4F8FC'
GRAY = '666666'
LIGHT_GRAY = 'E1E5E9'
VERY_LIGHT = 'F7F8FA'
BLACK = '000000'
WHITE = 'FFFFFF'
GREEN = '2F6B4F'
RED = '8A3030'

FONT_REG = '/System/Library/Fonts/Supplemental/Arial.ttf'
FONT_BOLD = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), color)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color=LIGHT_GRAY, size='6') -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = borders.find(qn(f'w:{edge}'))
        if tag is None:
            tag = OxmlElement(f'w:{edge}')
            borders.append(tag)
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), size)
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_run_font(run, name='Arial', size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_external_hyperlink(paragraph, text: str, url: str):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1F5A8A')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rPr.append(rFonts)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    hyperlink.append(r)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_internal_hyperlink(paragraph, text: str, anchor: str, size=10):
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1F5A8A')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size * 2))
    rPr.append(sz)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Arial')
    rFonts.set(qn('w:hAnsi'), 'Arial')
    rPr.append(rFonts)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    hyperlink.append(r)
    paragraph._p.append(hyperlink)


def add_bookmark(paragraph, name: str, bid: int) -> None:
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bid))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, text, end])


def set_picture_alt(inline_shape, alt: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set('descr', alt)
    doc_pr.set('title', alt[:120])


def font(size: int, bold=False):
    path = FONT_BOLD if bold else FONT_REG
    return ImageFont.truetype(path, size)


def pil_color(value: str) -> str:
    return value if value.startswith('#') else f'#{value}'


def rounded_box(draw, xy, fill, outline=BLUE, radius=24, width=4):
    draw.rounded_rectangle(xy, radius=radius, fill=pil_color(fill), outline=pil_color(outline), width=width)


def centered_multiline(draw, box, text, fnt, fill=BLACK, spacing=8):
    x1, y1, x2, y2 = box
    max_chars = max(10, int((x2-x1) / (fnt.size * 0.58)))
    lines = []
    for part in text.split('\n'):
        lines.extend(wrap(part, max_chars) or [''])
    block = '\n'.join(lines)
    bb = draw.multiline_textbbox((0, 0), block, font=fnt, spacing=spacing, align='center')
    w, h = bb[2]-bb[0], bb[3]-bb[1]
    draw.multiline_text(((x1+x2-w)/2, (y1+y2-h)/2), block, font=fnt, fill=pil_color(fill), spacing=spacing, align='center')


def arrow(draw, start, end, color=BLUE, width=6):
    draw.line([start, end], fill=pil_color(color), width=width)
    x2, y2 = end
    x1, y1 = start
    import math
    a = math.atan2(y2-y1, x2-x1)
    length = 18
    for delta in (2.55, -2.55):
        p = (x2 + length*math.cos(a+delta), y2 + length*math.sin(a+delta))
        draw.line([end, p], fill=pil_color(color), width=width)


def make_architecture(path: Path) -> None:
    im = Image.new('RGB', (1800, 1120), pil_color(WHITE))
    d = ImageDraw.Draw(im)
    title = font(46, True)
    h = font(32, True)
    body = font(27)
    d.text((70, 45), 'Bot first архитектура продукта', font=title, fill=pil_color(BLACK))
    bot = (80, 170, 810, 515)
    app = (990, 170, 1720, 515)
    core = (350, 680, 1450, 1020)
    rounded_box(d, bot, MID_BLUE)
    rounded_box(d, app, VERY_LIGHT)
    rounded_box(d, core, PALE_BLUE, outline=GREEN)
    centered_multiline(d, (100, 190, 790, 265), 'Основной канал MAX бот', h)
    centered_multiline(d, (120, 270, 770, 495), 'Вход и онбординг\nУточнение проблемы\nДиагностика и ответы\nКороткие уроки и практика\nСтатусы, напоминания, продолжение', body)
    centered_multiline(d, (1010, 190, 1700, 265), 'Мини приложение', h)
    centered_multiline(d, (1030, 270, 1680, 495), 'Сложные формы и визуализация\nКарта знаний и маршрут\nДлинные уроки\nФайлы и сертификаты\nНастройки и экспорт', body)
    centered_multiline(d, (380, 700, 1420, 775), 'Единый сервер и состояние', h, fill=GREEN)
    centered_multiline(d, (390, 790, 1410, 995), 'Fastify API и worker\nPostgreSQL, Redis, BullMQ, S3, ClamAV\nДиагностика, evidence, mastery, версии курсов\nYandex AI Studio и бюджет\nАудит, оператор и уведомления', body)
    arrow(d, (445, 515), (650, 680))
    arrow(d, (1355, 515), (1150, 680))
    arrow(d, (810, 342), (990, 342), color=GREEN)
    arrow(d, (990, 380), (810, 380), color=GREEN)
    im.save(path, quality=95)


def make_flow(path: Path) -> None:
    im = Image.new('RGB', (1800, 1050), pil_color(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 40), 'Полный пользовательский путь', font=font(46, True), fill=pil_color(BLACK))
    items = [
        ('Бот', 'Вход и проблема'), ('Диагностика', 'Бот или приложение'),
        ('Результат', 'Объяснение и цель'), ('Маршрут', 'Причины шагов'),
        ('Обучение', 'Урок и практика'), ('Проверка', 'Рубрика и evidence'),
        ('Перестройка', 'Новый маршрут'), ('Результат', 'Цель и сертификат')
    ]
    box_w, box_h = 360, 180
    coords = []
    for i, item in enumerate(items):
        row = i // 4
        col = i % 4 if row == 0 else 3 - (i % 4)
        x1 = 80 + col * 430
        y1 = 170 + row * 400
        xy = (x1, y1, x1+box_w, y1+box_h)
        coords.append(xy)
        rounded_box(d, xy, MID_BLUE if i in (0, 1, 2, 5, 7) else PALE_BLUE)
        centered_multiline(d, (x1+15, y1+18, x1+box_w-15, y1+78), item[0], font(30, True))
        centered_multiline(d, (x1+20, y1+78, x1+box_w-20, y1+box_h-15), item[1], font(25))
    for i in range(3):
        arrow(d, (coords[i][2], (coords[i][1]+coords[i][3])//2), (coords[i+1][0], (coords[i+1][1]+coords[i+1][3])//2))
    arrow(d, ((coords[3][0]+coords[3][2])//2, coords[3][3]), ((coords[4][0]+coords[4][2])//2, coords[4][1]))
    for i in range(4, 7):
        arrow(d, (coords[i][0], (coords[i][1]+coords[i][3])//2), (coords[i+1][2], (coords[i+1][1]+coords[i+1][3])//2))
    d.text((80, 965), 'Состояние хранится на сервере. Пользователь может продолжить путь из любого канала.', font=font(25), fill=pil_color(GRAY))
    im.save(path, quality=95)


def make_ai_pipeline(path: Path) -> None:
    im = Image.new('RGB', (1800, 920), pil_color(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 40), 'Контролируемое применение Yandex AI Studio', font=font(44, True), fill=pil_color(BLACK))
    labels = [
        'Версионный\nBuildRequest', 'Жёсткие\nфильтры', 'Разрешённые\nшаблоны',
        'Alice AI\nадаптация', 'JSON Schema\nвалидация', 'Candidate\ncourse', 'Атомарная\nпубликация'
    ]
    boxes = []
    for i, label in enumerate(labels):
        x1 = 55 + i * 247
        y1 = 280
        xy = (x1, y1, x1+205, y1+205)
        boxes.append(xy)
        fill = MID_BLUE if i == 3 else PALE_BLUE
        rounded_box(d, xy, fill, outline=RED if i == 3 else BLUE, radius=18, width=4)
        centered_multiline(d, xy, label, font(25, True if i in (0, 3, 6) else False))
        if i:
            arrow(d, (boxes[i-1][2], 382), (xy[0], 382), color=GREEN, width=5)
    d.text((70, 615), 'Решения остаются в коде', font=font(30, True), fill=pil_color(GREEN))
    d.text((70, 670), 'mastery, prerequisite, объективная проверка, права, сертификат, публикация', font=font(26), fill=pil_color(BLACK))
    d.text((70, 760), 'Модель помогает с текстом', font=font(30, True), fill=pil_color(RED))
    d.text((70, 815), 'классификация, уточнение, объяснение, адаптация урока, feedback по рубрике', font=font(26), fill=pil_color(BLACK))
    im.save(path, quality=95)


def make_roadmap(path: Path) -> None:
    im = Image.new('RGB', (1800, 1040), pil_color(WHITE))
    d = ImageDraw.Draw(im)
    d.text((70, 40), 'Порядок реализации', font=font(46, True), fill=pil_color(BLACK))
    phases = [
        ('0', 'Основа', 'Репозиторий\nконтракты'), ('1', 'Платформа', 'Инфраструктура\nмодель данных'),
        ('2', 'MAX', 'Бот\nauth и webhook'), ('3', 'Диагностика', 'Проблема\nвопросы и результат'),
        ('4', 'Курс', 'Каталог\nYandex AI и publish'), ('5', 'Обучение', 'Попытки\nevidence и mastery'),
        ('6', 'Полнота', 'Файлы\nоператор и данные'), ('7', 'Сдача', 'Тесты\nнагрузка и пакет')
    ]
    for i, (n, name, detail) in enumerate(phases):
        row = i // 4
        col = i % 4
        x1 = 70 + col * 430
        y1 = 180 + row * 390
        xy = (x1, y1, x1+360, y1+230)
        rounded_box(d, xy, MID_BLUE if i < 4 else PALE_BLUE)
        d.ellipse((x1+18, y1+18, x1+82, y1+82), fill=pil_color(BLUE))
        nb = d.textbbox((0,0), n, font=font(30, True))
        d.text((x1+50-(nb[2]-nb[0])/2, y1+50-(nb[3]-nb[1])/2), n, font=font(30, True), fill=pil_color(WHITE))
        centered_multiline(d, (x1+90, y1+15, x1+345, y1+90), name, font(29, True))
        centered_multiline(d, (x1+20, y1+100, x1+340, y1+215), detail, font(26))
        if col < 3:
            arrow(d, (x1+360, y1+115), (x1+430, y1+115), color=GREEN, width=5)
    d.text((70, 920), 'Сначала вертикальный сценарий на проверенных фикстурах. Затем модель и расширение охвата.', font=font(27), fill=pil_color(GRAY))
    im.save(path, quality=95)


ARCH = ASSETS / 'architecture_bot_first.png'
FLOW = ASSETS / 'user_flow.png'
AI_PIPE = ASSETS / 'ai_pipeline.png'
ROADMAP = ASSETS / 'roadmap.png'
make_architecture(ARCH)
make_flow(FLOW)
make_ai_pipeline(AI_PIPE)
make_roadmap(ROADMAP)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.72)
section.header_distance = Inches(0.28)
section.footer_distance = Inches(0.28)
section.different_first_page_header_footer = True

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = RGBColor.from_string(BLACK)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.12

for name, size, before, after in [('Title', 30, 0, 16), ('Subtitle', 15, 0, 12), ('Heading 1', 20, 14, 9), ('Heading 2', 15, 12, 6), ('Heading 3', 12, 9, 4)]:
    st = styles[name]
    st.font.name = 'Arial'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    st.font.size = Pt(size)
    st.font.bold = name != 'Subtitle'
    st.font.color.rgb = RGBColor.from_string(BLACK)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    if name == 'Heading 1':
        st.paragraph_format.page_break_before = True

# The stock Word Title style carries an accent-coloured bottom border.  The
# cover uses spacing and typography instead, so remove that inherited rule.
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)

prompt_label = styles.add_style('Prompt Label', 1)
prompt_label.font.name = 'Arial'
prompt_label._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
prompt_label._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
prompt_label.font.size = Pt(9)
prompt_label.font.bold = True
prompt_label.font.color.rgb = RGBColor.from_string(BLUE)
prompt_label.paragraph_format.space_before = Pt(7)
prompt_label.paragraph_format.space_after = Pt(3)
prompt_label.paragraph_format.keep_with_next = True

prompt_style = styles.add_style('Prompt Body', 1)
prompt_style.font.name = 'Consolas'
prompt_style._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
prompt_style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
prompt_style.font.size = Pt(8.7)
prompt_style.font.color.rgb = RGBColor.from_string('333333')
prompt_style.paragraph_format.left_indent = Inches(0.28)
prompt_style.paragraph_format.right_indent = Inches(0.08)
prompt_style.paragraph_format.space_after = Pt(9)
prompt_style.paragraph_format.line_spacing = 1.04

caption_style = styles['Caption']
caption_style.font.name = 'Arial'
caption_style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
caption_style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
caption_style.font.size = Pt(9)
caption_style.font.italic = True
caption_style.font.color.rgb = RGBColor.from_string(GRAY)

header = section.header
hp = header.paragraphs[0]
hp.text = 'Образовательная платформа MAX  Книга разработчика'
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    set_run_font(r, size=8.5, color=GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run('Страница ')
set_run_font(fr, size=8.5, color=GRAY)
add_page_number(fp)

bookmark_counter = 10
toc_entries: list[tuple[int, str, str]] = []


def heading(text: str, level: int = 1, toc=True):
    global bookmark_counter
    p = doc.add_paragraph(text, style=f'Heading {level}')
    anchor = f'section_{bookmark_counter:03d}'
    add_bookmark(p, anchor, bookmark_counter)
    bookmark_counter += 1
    if toc:
        toc_entries.append((level, text, anchor))
    if level == 1:
        back = doc.add_paragraph()
        back.paragraph_format.space_after = Pt(4)
        add_internal_hyperlink(back, 'К содержанию', 'toc_anchor', size=8)
    return p


def body(text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def bullets(items: Iterable[str], level=0):
    for item in items:
        p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
        p.paragraph_format.space_after = Pt(2.5)
        r = p.add_run(item)
        set_run_font(r)


def numbered(items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def table(headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = 1
    set_table_borders(t)
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        set_cell_fill(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, size=9, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[i])
    repeat_table_header(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, value in enumerate(row):
            cells[ci].text = value
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[ci])
            if ri % 2:
                set_cell_fill(cells[ci], PALE_BLUE)
            if widths:
                cells[ci].width = Inches(widths[ci])
            for p in cells[ci].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def figure(path: Path, caption: str, alt: str, width=6.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt(shape, alt)
    cp = doc.add_paragraph(caption, style='Caption')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def source_line(label: str, value: str, url: str | None = None):
    p = doc.add_paragraph()
    r = p.add_run(label + '  ')
    set_run_font(r, bold=True)
    if url:
        add_external_hyperlink(p, value, url)
    else:
        r = p.add_run(value)
        set_run_font(r, size=9.5, color=GRAY)


PROMPT_SPECIFIC_VERIFICATIONS = {
    '00': ['Проверь, что каждая строка карты ссылается на существующий источник и имеет способ проверки.', 'Сверь количество FR, типов запроса и конфликтов с исходными документами; пустые или противоречивые записи исправь.'],
    '01': ['Сравни каждый путь манифеста с legacy/vibework_reference и найди отсутствующие или лишние файлы.', 'Запусти поиск секретов, .env, БД, кэшей и персональных данных внутри выборки; всё найденное удали из выборки и повтори поиск.'],
    '02': ['Проверь уникальность записей, допустимые статусы и наличие owner, fallback и impact.', 'Для каждого отсутствующего ключа смоделируй запуск без него и убедись, что приложение остаётся работоспособным в mock или disabled режиме.'],
    '03': ['Выполни чистую установку зависимостей, затем build, typecheck, lint и test из корня.', 'Запусти по одной минимальной команде каждого workspace и проверь отсутствие циклических или неразрешённых импортов.'],
    '04': ['Провалидируй все примеры по JSON Schema и все 17 путей OpenAPI.', 'Подай заведомо неверные payload и убедись, что runtime validation возвращает стабильную ошибку, а расхождения попадают в CONTRACT_DEVIATIONS.'],
    '05': ['Прогони одинаковую ошибку через API, worker и тестовый adapter и сравни код, trace id и безопасное сообщение.', 'Проверь повтор с тем же idempotency key и разные branded id, включая невалидный последовательный идентификатор.'],
    '06': ['Выполни docker compose config, подними стек, дождись health checks и корректно останови его.', 'Перезапусти контейнеры с пустыми volumes и проверь работу без внешних ключей и без бесконечных рестартов.'],
    '07': ['Примени миграции к пустой БД, откати допустимый шаг и примени повторно.', 'Запусти миграции второй раз и проверь идемпотентность, ограничения внешних ключей и запрет невалидных переходов.'],
    '08': ['Проверь успешный job, повтор после временной ошибки, timeout, dead letter и ручной retry.', 'Отправь дубликаты одного job и убедись, что доменный результат и побочные эффекты создаются один раз.'],
    '09': ['Смоделируй rollback доменной транзакции и проверь отсутствие сиротского outbox события.', 'Доставь одно outbox событие дважды и проверь дедупликацию потребителя и полный audit trail.'],
    '10': ['Отправь валидный, неверно подписанный, повторный и повреждённый webhook.', 'Замерь ответ endpoint и проверь, что подтверждение укладывается в лимит, а тяжёлая работа выполняется только worker.'],
    '11': ['На mock transport проверь 2xx, 429, 5xx, timeout и malformed response.', 'Проверь безопасный retry только идемпотентных операций и отсутствие токена или персональных данных в логах.'],
    '12': ['Пройди все разрешённые переходы BotConversationState и запрети недопустимые.', 'Проверь устаревший callback, сообщение не по порядку и повтор команды без повреждения доменного состояния.'],
    '13': ['Пройди новый вход, согласие, отказ, возрастное ограничение и повторный вход.', 'Проверь, что отказ не создаёт профиль, а прерванный разрешённый сценарий продолжается с сохранённого состояния.'],
    '14': ['Проверь понятный запрос, неоднозначный запрос, изменение формулировки и полный перезапуск.', 'Убедись, что без явного подтверждения не создаются ProblemVersion и следующие диагностические шаги.'],
    '15': ['Пройди single, multiple, skip unknown, pause, resume и досрочное завершение.', 'Проверь верхний лимит вопросов, отсутствие повторов и сохранение каждого ответа до следующего сообщения.'],
    '16': ['Проверь результаты known gap, unknown и partial evidence, затем подтверждение и исправление цели.', 'Убедись, что бот не сообщает ложную точность и не требует mini app для продолжения основного пути.'],
    '17': ['Пройди короткий урок целиком в боте, включая возврат после паузы.', 'Проверь длинный или сложный материал: бот должен дать полезное резюме и осмысленную ссылку на mini app, сохранив текущий шаг.'],
    '18': ['Проверь правильный, неправильный и повторный ответ с привязанной rubric version.', 'Убедись, что feedback воспроизводим, попытка неизменяема и открытие mini app не является обязательным.'],
    '19': ['Доставь одно событие уведомления дважды, с отключёнными уведомлениями и после смены маршрута.', 'Проверь, что кнопка продолжения ведёт в актуальный шаг, а сообщение не раскрывает чувствительные данные.'],
    '20': ['Отправь неизвестную команду, устаревший callback и сообщение при отсутствующем состоянии.', 'Убедись, что пользователь получает восстановимый следующий шаг, а stack trace и данные провайдера не показываются.'],
    '21': ['Запусти приложение в обычном браузере с mock bridge и внутри MAX с реальным bridge contract.', 'Проверь loading, empty, error и resume state на мобильной ширине без блокирующей ошибки bridge.'],
    '22': ['Проверь валидный, изменённый, просроченный и повторно использованный initData.', 'Проверь CSRF, cookie flags, logout и отсутствие bot token в браузере и сетевых ответах.'],
    '23': ['Отправь минимально допустимую форму, невалидные поля и конфликт версии.', 'Проверь autosave, повторную отправку и отсутствие обязательного карьерного теста для учебной проблемы.'],
    '24': ['Проверь single, multiple, skip unknown, клавиатурное управление и восстановление после reload.', 'Убедись, что повторный submit не создаёт второй ответ и адаптивная прогрессия выбирает ожидаемый вопрос.'],
    '25': ['Проверь полный, частичный и недостаточный evidence набор.', 'Убедись, что интерфейс показывает ограничения и provenance, не рисует проценты без методики и допускает корректировку цели.'],
    '26': ['Проверь актуальную пару route course, устаревшую версию и атомарное переключение на новую.', 'Убедись, что причины шага и заблокированные предпосылки видимы, а старый маршрут остаётся только в истории.'],
    '27': ['Пройди урок с текстовым ответом, файлом, ошибкой загрузки, review и повторной попыткой.', 'Проверь связь Attempt и ReviewVersion, отображение rubric и сохранение данных после перезагрузки.'],
    '28': ['Проверь незавершённую и завершённую цель, доступность сертификата, экспорт и удаление.', 'Проверь восстановление настроек, deep link из MAX и понятное подтверждение необратимого действия.'],
    '29': ['Прогони интерфейс на 360, 390 и 1280 пикселях, с клавиатурой и screen reader.', 'Проверь contrast, focus order, target size, reduced motion и отсутствие горизонтального скролла.'],
    '30': ['Импортируй валидный каталог, дубликат версии, неизвестную ссылку и запись без provenance.', 'Проверь повторяемость импорта, запрет DRAFT и ARCHIVED в runtime и отсутствие автоматической публикации.'],
    '31': ['Прогони fixtures всех четырёх типов запроса и неоднозначный случай.', 'Повтори одинаковый ввод несколько раз и проверь детерминированный fallback, сохранённую версию prompt и объяснение выбора.'],
    '32': ['Построй план для каждого типа запроса, неизвестной проблемы и ограниченного времени.', 'Проверь hard limit, отсутствие запрещённых вопросов и остановку при достаточном evidence.'],
    '33': ['Смоделируй несколько наборов evidence и проверь ожидаемый следующий вопрос.', 'Проверь отсутствие повторов, зависания на одной области и выход по лимиту или stopping condition.'],
    '34': ['Собери результаты known gap, unknown и partial и сравни с ожидаемыми fixtures.', 'Проверь воспроизводимость результата, ссылки answer submission и запрет раскрытия правильных ответов.'],
    '35': ['Прогони fake provider для success, retryable error, permanent error и timeout.', 'Проверь одинаковый нормализованный ответ всех адаптеров и отсутствие сетевой зависимости в domain тестах.'],
    '36': ['На recorded fixtures проверь обязательные заголовки, auth, 429, 5xx, timeout и malformed JSON.', 'Убедись, что retry ограничен, logging и store отключаются настройками, а ключ не попадает в клиент или логи.'],
    '37': ['Проверь запрос ниже бюджета, запрос выше лимита и превышение после фактического usage.', 'Сверь расчёт цены по токенам, атомарность ledger и отсутствие повторного списания при retry.'],
    '38': ['Проверь каждый prompt template на валидном ответе, невалидном JSON и лишних полях.', 'Добавь prompt injection fixtures и убедись, что trusted instructions, schema и разрешённый контекст нельзя переопределить пользовательским вводом.'],
    '39': ['Запусти benchmark на фиксированном наборе минимум дважды и сравни стабильность метрик.', 'Проверь пороги schema validity, domain validity, latency и cost; выбор победителя должен следовать данным отчёта.'],
    '40': ['Повтори build с одинаковым snapshot и проверь тот же hash и результат.', 'Проверь цикл графа, отсутствующую компетенцию, запрещённый источник и изменение активной версии во время build.'],
    '41': ['Проверь возрастные, временные, source и prerequisite ограничения на граничных значениях.', 'Убедись, что tie breaking детерминирован и запрещённый материал не попадает в контекст модели.'],
    '42': ['Проверь успешную адаптацию, timeout, невалидный structured output и отключённый AI.', 'Сравни source truth до и после: ответы, компетенции и обязательные инструкции не должны изменяться.'],
    '43': ['Запусти два конкурирующих build и публикацию устаревшего candidate.', 'Проверь rollback при ошибке, атомарное переключение route course и единственное событие published.'],
    '44': ['Создай несколько попыток и ReviewVersion, затем повторно отправь одинаковый запрос.', 'Проверь неизменяемость истории, ссылку на rubric version и очистку AI feedback от непроверенных утверждений.'],
    '45': ['Проверь правило двух PASS разных family, повтор той же family и один FAIL.', 'Проверь revoked, expired и disputed evidence, а также порог 90 дней без зависимости от AI.'],
    '46': ['Проверь replan по новому evidence, изменению цели и незначимому событию.', 'Убедись, что старая версия доступна в истории, активная пара переключается атомарно и причина записана в audit.'],
    '47': ['Пройди создание, рассмотрение, решение и повторное открытие спора.', 'Проверь RBAC, SLA, конкурентные решения и запрет автоматического изменения mastery до решения оператора.'],
    '48': ['Одинаковым contract suite проверь local и S3 adapter: put, get, delete и signed URL.', 'Проверь owner authorization, path traversal, MIME mismatch, лимит размера и недоступное хранилище.'],
    '49': ['Загрузи чистый файл, EICAR, неверный MIME, превышение размера и файл при недоступном scanner.', 'Проверь quarantine state, запрет скачивания до clean и audit каждого перехода.'],
    '50': ['Проверь допустимую и недопустимую выдачу, повторный запрос и отозванное evidence.', 'Проверь целостность PDF, стабильный certificate id и невозможность создать сертификат только через UI.'],
    '51': ['Проверь полный экспорт, запрос удаления, grace period, отмену и окончательную очистку.', 'Повтори jobs и проверь идемпотентность, audit trail, retention исключения и удаление связанных файлов.'],
    '52': ['Для каждой роли проверь разрешённые и запрещённые действия и прямой вызов API без UI.', 'Проверь аудит просмотра и изменения, отсутствие персональных данных в списках и безопасную пагинацию.'],
    '52A': ['Проверь, что каждый файл из SOURCE_MANIFEST.md имеет строку с переносимым поведением, целевым модулем, заменой технологии и набором тестов.', 'Автоматически проверь целевой репозиторий на запрещённые runtime зависимости Python, FastAPI и SQLite; найденные совпадения классифицируй как документацию, fixture либо нарушение.'],
    '53': ['Конвертируй каталог дважды и сравни checksum, версии и количество записей.', 'Проверь битую ссылку, неизвестную лицензию, отсутствие provenance и запрет автоматического статуса PUBLISHED.'],
    '54': ['Прогони четыре типа пользовательского запроса и убедись, что карьерная ветка включается только для выбора направления.', 'Проверь отсутствие старых психологических ярлыков, утечки ответов и запуска карьерного опроса в учебном сценарии.'],
    '55': ['Проверь mock HH и MTS для успеха, пустого ответа, rate limit и недоступности.', 'Убедись, что внешние запросы не выполняются до подтверждённой цели, а возрастные и юридические ограничения соблюдаются.'],
    '56': ['Пройди симулятор для нескольких сценариев и сравни оценку с фиксированной rubric.', 'Проверь retry, resume и то, что прохождение само по себе не создаёт независимое mastery evidence без review.'],
    '56A': ['Одинаковым набором fixture проверь старые негативные случаи и новую TypeScript реализацию: размер, MIME, magic bytes, owner, quarantine и EICAR.', 'Проверь сборку и запуск API и worker в окружении без Python; ни один путь загрузки, сканирования или скачивания не должен запускать Python subprocess или читать SQLite.'],
    '57': ['Прогони success и документированные ошибки для всех 17 путей OpenAPI.', 'Проверь ApiError, соответствие generated types и fixtures, а также machine readable отчёт без молча пропущенных путей.'],
    '58': ['Запусти полный набор unit fixtures для инвариантов known gap, mastery, версии и атомарной публикации.', 'Временно внеси контролируемую мутацию одного правила и убедись, что соответствующий тест падает, затем верни код и повтори suite.'],
    '59': ['Смоделируй падение worker, повтор webhook, timeout Yandex, 503 S3 и недоступный Redis.', 'Проверь отсутствие дублей, зависших состояний и потери события; после восстановления система должна завершить или явно пометить операцию.'],
    '60': ['Пройди полный сценарий от первого сообщения до повторной попытки только через бот.', 'Проверь restart, resume и фоновые уведомления; ни на одном обязательном шаге единственным действием не должно быть открытие mini app.'],
    '61': ['Начни в боте, продолжи в mini app и вернись в бот на нескольких точках сценария.', 'Проверь общий version state, deep links, mobile и web режим, session expiry и отсутствие аналитики с персональными данными.'],
    '62': ['Запусти профиль нагрузки из ТЗ и сохрани p50, p95, error rate, CPU и память.', 'Проверь деградацию при перегрузе, отсутствие duplicate jobs и автоматизированное сравнение с порогами.'],
    '63': ['Запусти secret scan, dependency audit, SAST и целевые abuse tests.', 'Проверь подделку initData, replay webhook, IDOR, SSRF, prompt injection, path traversal и утечки в логах; high severity должен блокировать release.'],
    '64': ['Повтори запуск по README из чистого checkout без локальных незадокументированных файлов.', 'Проверь все health checks, migrations, seed, mocks, список переменных и отсутствие секретов или абсолютных путей разработчика в пакете.'],
    '65': ['Трижды воспроизведи демонстрационный сценарий с чистыми demo данными и зафиксируй результат.', 'Проверь offline или mock fallback, редактирование скриншотов и логов, соответствие доказательств критериям конкурса и отсутствие вымышленных метрик.'],
}


LEGACY_PORTING_RULES = {
    '30': [
        'Бери из learning_catalog.json, learning_paths.json и services/learning только данные, ограничения и ожидаемые случаи. Импортёр и runtime реализуй на TypeScript в packages/content и apps/worker.',
        'Старые Python loader, SQLite queries и неявные словари состояния не подключай. Целевое хранение — PostgreSQL через репозитории и миграции; валидация — JSON Schema/Ajv.',
        'Перенеси полезные Python тестовые случаи в TypeScript fixtures и докажи эквивалентность результата на фиксированном наборе данных.'
    ],
    '31': [
        'Используй user_pain_mapping.py как спецификацию поведения и источник примеров, но перепиши классификатор чистыми TypeScript функциями в packages/domain.',
        'Не запускай Python как subprocess или sidecar, не импортируй FastAPI и не делай старый LLM клиент обязательным для классификации.',
        'Перенеси таблицы соответствий и граничные случаи в версионные TypeScript fixtures; детерминированный fallback должен работать без сети.'
    ],
    '32': [
        'Используй assessment_routing.py и assessment_modules.py только как источник правил и вопросов. Планировщик диагностики реализуй в TypeScript как доменный сервис.',
        'Замени Python dataclass и глобальное состояние типизированными immutable структурами; версии плана и ответы сохраняй в PostgreSQL.',
        'Перепиши релевантные Python тесты на текущий test runner TypeScript и добавь проверку отсутствия вызовов старого backend.'
    ],
    '33': [
        'Перенеси алгоритмическую идею выбора следующего вопроса, а не Python реализацию. Новый selector должен быть чистой TypeScript функцией с явными входами, clock и seed.',
        'Не используй SQLite, module-level cache или mutable dict из VibeWork; состояние сессии читай и записывай через PostgreSQL repository.',
        'Сравни новую реализацию с зафиксированными legacy fixtures, затем расширь их случаями нового контракта MAX.'
    ],
    '34': [
        'Используй diagnostics.py и profile_analysis_context.py как поведенческие доноры. DiagnosticResult и DiagnosticContext собирай в TypeScript по актуальным JSON Schema.',
        'Не переносись на старые Pydantic модели, FastAPI response и SQLite. Все ссылки на evidence должны быть типизированы и храниться в PostgreSQL.',
        'Старые expected cases преобразуй в TypeScript fixtures; расхождения с новым ТЗ разрешай в пользу контрактов комплекта 4.0 и фиксируй в документации.'
    ],
    '38': [
        'Из llm_prompts.py разрешено брать только смысл шаблонов, примеры и fallback-тексты. Перепиши их как версионные TypeScript templates с JSON Schema structured output.',
        'Старый облачный LLM клиент и его формат ответа не переносить. Все вызовы выполняются через общий TypeScript AI provider port и Yandex AI adapter.',
        'Для каждого перенесённого шаблона добавь snapshot/contract tests, injection fixtures и тест режима без внешнего ключа.'
    ],
    '42': [
        'Используй learning services, action_plans.py и рекомендации VibeWork только как источник правил персонализации; исполняемый код перепиши на TypeScript.',
        'Детерминированный выбор материалов реализуй в domain/content, а генеративную адаптацию — через Yandex provider adapter и BullMQ job. Не вызывай Python или старый LLM клиент.',
        'Проверь, что выключенный AI оставляет рабочий TypeScript fallback и не меняет утверждённые ответы, рубрики и источники.'
    ],
    '48': [
        'profile_files.py используй только как перечень негативных сценариев. Контракт FileStorage, local adapter и S3 adapter реализуй на TypeScript для Fastify.',
        'Не переноси UploadFile, Pillow, tempfile и локальные абсолютные пути. Используй потоковую загрузку, magic-byte проверку, ограничения размера и server-side owner authorization.',
        'Перепиши legacy проверки в единый TypeScript contract suite, одинаковый для local и S3 реализаций.'
    ],
    '49': [
        'Из profile_files.py перенеси только правила валидации и тестовые случаи. Оркестрацию карантина реализуй TypeScript worker на BullMQ с ClamAV adapter.',
        'Не используй Python subprocess, Pillow и SQLite. Метаданные и статусы храни в PostgreSQL, содержимое — через FileStorage/S3.',
        'Добавь TypeScript integration tests для clean, infected, MIME mismatch, scanner timeout и повторной задачи.'
    ],
    '52A': [
        'Считай Python код VibeWork спецификацией поведения, а не зависимостью нового приложения. Никакой Python файл не должен загружаться, исполняться или вызываться новым runtime.',
        'Зафиксируй обязательные замены: Python/FastAPI на Node.js/TypeScript/Fastify; SQLite на PostgreSQL; requests на типизированные HTTP adapters; старый LLM client на Yandex AI provider; долгие операции на BullMQ/Redis; файловые операции на S3/ClamAV.',
        'Для каждого переносимого элемента укажи, какие fixtures и тесты сохраняют поведение, а какие детали старой реализации намеренно отбрасываются.'
    ],
    '53': [
        'Считывай исходные данные только из legacy/vibework_reference/website/data. Конвертер создай как TypeScript CLI в packages/content или apps/worker.',
        'Не подключай Python loaders и SQLite. Результат — версионный канонический JSON и/или идемпотентный импорт в PostgreSQL через целевой repository.',
        'Перенеси проверки ссылочной целостности и данных в TypeScript tests; runtime нового продукта не должен зависеть от папки legacy.'
    ],
    '54': [
        'career_navigator.py, role_confirmation.py и career_analysis файлы используй только как источник правил, терминов и fixtures. Реализуй политику выбора направления в packages/domain на TypeScript.',
        'Не переноси FastAPI routes, Pydantic models, SQLite access и глобальный профиль пользователя. Состояние хранится через PostgreSQL repositories и общие контракты MAX.',
        'Перепиши полезные тесты на TypeScript и отдельно докажи, что учебный запрос не запускает карьерную ветку.'
    ],
    '55': [
        'Из hh_filter.py, job_search.py и mts_match.py бери правила фильтрации и role mappings. Реальные интеграции реализуй TypeScript adapters с типизированными DTO.',
        'Не переноси Python requests, старые config globals и сетевые вызовы из роутов. Используй серверный HTTP client, env validation, timeout, rate limit и Redis cache.',
        'Перепиши legacy fixtures на TypeScript; mock и реальные providers должны проходить один contract suite и не попадать в клиентский bundle.'
    ],
    '56': [
        'workday_simulator.py и simulator_progress.py используй как источник сценариев и переходов. Домен симулятора реализуй на TypeScript с versioned scenario и rubric.',
        'Не переноси Python dict, SQLite progress и неявные module globals. Попытки, события и resume state сохраняй в PostgreSQL.',
        'Перепиши сценарные тесты на TypeScript; bot и mini app должны использовать одну серверную модель и одинаковый результат оценки.'
    ],
    '56A': [
        'Сопоставь каждый полезный негативный случай profile_files.py с TypeScript тестом и текущим контрактом FileStorage; сам Python модуль не копируй в runtime.',
        'Целевая цепочка: Fastify stream — quarantine — S3/local adapter — BullMQ worker — ClamAV — PostgreSQL status. Все связи реализуй на TypeScript.',
        'Удаляй временные файлы безопасно, не доверяй имени и Content-Type клиента и не разрешай download до статуса clean.'
    ],
}


def prompt_block(pid: str, title: str, goal: str, scope: list[str], checks: list[str], inputs: list[str] | None = None, notes: list[str] | None = None):
    heading(f'Промт {pid} {title}', 2)
    label = doc.add_paragraph('КОПИРОВАТЬ В CODEX ЦЕЛИКОМ', style='Prompt Label')
    label.paragraph_format.keep_with_next = True
    lines = [
        'Работай в репозитории /Users/inkoromi21/Documents/VibeWorkMax. Перед изменениями прочитай AGENTS.md, README и связанные исходные файлы. Для старого проекта сначала используй только legacy/vibework_reference и SOURCE_MANIFEST.md; полный /Users/inkoromi21/Documents/VibeWork открывай только при доказанной нехватке выборки и объясни это в отчёте. Сохраняй пользовательские изменения, не переноси секреты и персональные данные, не меняй исходные документы в Downloads. Реализуй только ограниченную задачу ниже. Если репозиторий ещё не создан, подготовь только те минимальные файлы, которые нужны этой задаче.',
        '',
        f'Цель: {goal}',
        '',
        'Что сделать:'
    ]
    for i, item in enumerate(scope, 1):
        lines.append(f'{i}. {item}')
    if inputs:
        lines += ['', 'Нужные внешние данные:']
        for item in inputs:
            lines.append(f'- {item}')
        lines.append('Если данных нет, не выдумывай их: добавь переменные в .env.example, интерфейс адаптера, безопасный mock или feature flag и запись в docs/MISSING_INPUTS.md. Остальную работу продолжай.')
    if notes:
        lines += ['', 'Ограничения:']
        for item in notes:
            lines.append(f'- {item}')
    if pid in LEGACY_PORTING_RULES:
        lines += ['', 'Правила переписывания на целевой стек:']
        for item in LEGACY_PORTING_RULES[pid]:
            lines.append(f'- {item}')
    lines += ['', 'Критерии готовности:']
    for item in checks:
        lines.append(f'- {item}')
    lines += ['', 'Проверки работоспособности:']
    verification_steps = [
        'Запусти форматирование, lint и проверку типов только для затронутых пакетов; новые ошибки и предупреждения устрани.',
        'Добавь или обнови автоматические тесты для позитивного, граничного и ошибочного сценариев изменения.',
        *PROMPT_SPECIFIC_VERIFICATIONS[pid],
        'Запусти связанные существующие тесты и минимальный smoke сценарий затронутого сервиса, worker или интерфейса.',
        'Проверь логи, БД, очередь и внешние побочные эффекты: не должно быть необработанных ошибок, дублей, секретов и зависших состояний.',
    ]
    for i, item in enumerate(verification_steps, 1):
        lines.append(f'{i}. {item}')
    lines += [
        '',
        'Обязательный цикл исправления: выполни весь набор проверок выше. Если хотя бы одна проверка падает из-за внесённых изменений, найди причину, исправь её и повтори весь затронутый набор до успешного результата. Не маскируй ошибку ослаблением assertion, отключением теста, добавлением необоснованного retry или исключением файла из проверки. Если сбой внешний или существовал до задачи, докажи это минимальным воспроизведением, запиши блокер и оставь безопасный fallback. В конце перечисли изменённые файлы, все команды проверки, фактические результаты, исправленные проблемы, оставшиеся риски и отсутствующие входные данные.'
    ]
    p = doc.add_paragraph('\n'.join(lines), style='Prompt Body')
    p.paragraph_format.widow_control = True


# Cover
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(110)
r = p.add_run('Образовательная платформа MAX')
set_run_font(r, size=30, bold=True)
p2 = doc.add_paragraph(style='Title')
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run('Книга разработчика')
set_run_font(r, size=28, bold=True)

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.paragraph_format.space_before = Pt(18)
r = sp.add_run('Архитектура бот и мини приложение\nПеренос решений из VibeWork\nYandex AI Studio\nПолный план реализации и библиотека промтов для Codex')
set_run_font(r, size=14, color=GRAY)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(95)
r = meta.add_run('Рабочая редакция 25 сентября 2026 года')
set_run_font(r, size=11, color=GRAY)

doc.add_page_break()

# TOC marker
toc_title = doc.add_paragraph('Содержание', style='Heading 1')
add_bookmark(toc_title, 'toc_anchor', 1)
toc_marker = doc.add_paragraph('[[STATIC_TOC_ENTRIES]]')
toc_end = doc.add_paragraph()
toc_end.add_run().add_break(WD_BREAK.PAGE)

heading('Как пользоваться этой книгой', 1)
body('Эта книга одновременно является архитектурным отчётом, дорожной картой и библиотекой готовых заданий для Codex. Сначала прочитайте главы о границах продукта и конфликтах требований. Затем выполняйте промты последовательно. Каждый промт ограничен одной проверяемой задачей и рассчитан на отдельный рабочий цикл.')
bullets([
    'Выполняйте промты по порядку, если проект создаётся с нуля.',
    'Перед каждым промтом проверяйте, закрыты ли его зависимости из предыдущих задач.',
    'Не объединяйте несколько промтов в одну огромную задачу без необходимости.',
    'Если отсутствует ключ, домен или контент, Codex должен создать интерфейс, mock и запись о блокере, но не придумывать секрет.',
    'После каждой фазы делайте отдельный стабильный commit и запускайте проверки этой фазы.',
    'Документы конкурса и тимлида являются источником требований. VibeWork является только донором идей и переносимой логики.'
])

heading('Основной вывод', 1)
body('Новый продукт нельзя строить как существующий сайт VibeWork, спрятанный за кнопкой MAX. Основным решением должен быть бот с самостоятельным полезным сценарием. Мини приложение расширяет возможности бота там, где нужны сложные формы, визуализация, длинные материалы, работа с файлами и управление данными. Оба канала работают с одним сервером и одним состоянием пользователя.')
body('Разработка должна идти contracts first. Сначала создаётся полностью работающий вертикальный сценарий на проверенных фикстурах, затем подключается Yandex AI Studio, после чего добавляются остальные типы запросов и необязательные модули VibeWork.')
figure(ARCH, 'Рисунок 1  Распределение ответственности в bot first архитектуре', 'Схема bot first архитектуры с ботом, мини приложением и единым сервером')

heading('Иерархия требований', 1)
body('При конфликте источников используйте следующий порядок приоритета.')
numbered([
    'Законодательство и обязательные правила платформы MAX.',
    'Обязательные условия конкурса и критерии допуска.',
    'Техническое задание тимлида и дополнение к нему.',
    'Существующий код и продуктовые решения VibeWork.'
])
body('Конкурс разрешает разные языки программирования. Техническое задание тимлида требует React и TypeScript на клиенте, Node.js и Fastify на сервере, PostgreSQL, Redis, BullMQ, S3 и ClamAV. Если тимлид не дал отдельного разрешения, внутреннее более строгое требование следует считать обязательным.')

heading('Карта конфликтов и решений', 1)
table(
    ['Область', 'Конфликт', 'Принятое решение'],
    [
        ['Роль бота', 'Мини приложение может стать фактическим продуктом, а бот только кнопкой запуска', 'Бот получает собственную диагностику, обучение, практику, статусы и продолжение сценария'],
        ['Стек', 'Конкурс допускает другие языки, VibeWork написан на Python', 'Целевой runtime делается на TypeScript; Python используется как источник алгоритмов'],
        ['Масштаб', 'Тимлид требует все функции первой стадии, конкурс советует узкий MVP', 'Общий движок поддерживает четыре запроса, основной демо контент посвящён бесконечному циклу'],
        ['Профориентация', 'VibeWork начинает с длинного профиля и тестов', 'Новый путь начинается с проблемы; карьерные тесты включаются только при выборе направления'],
        ['LLM', 'VibeWork доверяет модели генерацию плана', 'Модель создаёт только candidate и текст; решения принимают правила и валидаторы'],
        ['Прогресс', 'В VibeWork шаг можно отметить выполненным', 'Освоение подтверждается evidence и независимыми попытками'],
        ['Демо данные', 'Комплект содержит синтетические фикстуры', 'Они явно маркируются и не выдаются за экспертно подтверждённые данные'],
        ['API', 'Конкурс не даёт бонус за собственный API', 'API всё равно реализуется по контрактам тимлида'],
    ], [1.05, 2.65, 3.35]
)

heading('Что должен уметь основной бот', 1)
body('Бот должен позволять пройти значительную часть главного сценария без открытия мини приложения. Это делает его самостоятельным продуктовым каналом и снижает риск несоответствия условию конкурса.')
table(
    ['Группа', 'Функции бота', 'Когда нужен переход в приложение'],
    [
        ['Вход', 'Приветствие, согласия, возрастная проверка, выбор начать или продолжить', 'Подробные юридические тексты и настройки согласий'],
        ['Проблема', 'Свободный текст, уточнение, подтверждение интерпретации, выбор типа запроса', 'Редактирование большого профиля или сложных ограничений'],
        ['Диагностика', 'Один вопрос на сообщение, кнопки ответа, не знаю, пропустить, пауза', 'Практическое задание с кодом, большой схемой или файлом'],
        ['Результат', 'Кратко известное, пробел, неизвестное, первая рекомендация, подтверждение цели', 'Карта знаний, подробные причины и сравнение вариантов'],
        ['Обучение', 'Короткий урок, микро пример, контрольный вопрос, следующая практика', 'Длинный урок, интерактив, сложная визуализация'],
        ['Практика', 'Текстовый ответ, выбор варианта, короткое объяснение, повторная попытка', 'Файлы, развёрнутый ответ, просмотр рубрики и истории версий'],
        ['Продолжение', 'Статус, напоминание, кнопка продолжить с текущего шага', 'Полная навигация по курсу'],
        ['Итог', 'Сообщение о достижении цели, ссылка на сертификат, следующая возможность', 'Скачивание PDF, экспорт и управление данными'],
    ], [1.05, 3.35, 2.65]
)

heading('Что должно быть в мини приложении', 1)
bullets([
    'Редактирование профиля и ограничений пользователя.',
    'Визуальная диагностика с прогрессом и сложными типами вопросов.',
    'Карта знаний и объяснение результата.',
    'Полный маршрут обучения и причины включения шагов.',
    'Длинные уроки, вложенные материалы и интерактивные задания.',
    'Отправка и скачивание файлов.',
    'История попыток, подробная рубрика и оспаривание результата.',
    'Прогресс по компетенциям и версии маршрута.',
    'Сертификаты, следующие возможности, экспорт и удаление данных.',
    'Состояния загрузки, паузы, офлайн режима и восстановления.'
])
figure(FLOW, 'Рисунок 2  Пользовательский путь между ботом приложением и сервером', 'Схема полного пользовательского пути от бота до сертификата')

heading('Как каналы работают с одним состоянием', 1)
numbered([
    'Webhook бота проходит проверку секрета и дедупликацию.',
    'Сервер сохраняет событие и outbox в короткой транзакции и отвечает MAX.',
    'Диалоговый обработчик переводит пользователя между теми же доменными состояниями, что и mini app.',
    'Mini app передаёт серверу исходную строку initData, а сервер проверяет подпись и срок действия.',
    'Кнопка из бота содержит только непрозрачный start parameter, ведущий к разрешённому действию.',
    'Любое действие в одном канале сразу отражается в другом, потому что состояние хранится только на сервере.',
    'Долгая генерация выполняется worker процессом. Бот сообщает, когда результат готов.',
    'Повторное событие или повторный callback не создаёт вторую попытку, курс или уведомление.'
])

heading('Целевая техническая архитектура', 1)
table(
    ['Компонент', 'Технология', 'Ответственность'],
    [
        ['Bot transport', 'MAX Bot API', 'Webhook, callbacks, сообщения, клавиатуры, deep links и уведомления'],
        ['Mini app', 'React TypeScript Vite MAX UI MAX Bridge', 'Сложные экраны, визуализация, файлы и настройки'],
        ['API', 'Node.js TypeScript Fastify', 'Контракты, auth, доменные команды и чтение состояния'],
        ['Worker', 'Тот же TypeScript codebase', 'Yandex AI, проверка файлов, PDF, уведомления и фоновые задания'],
        ['Database', 'PostgreSQL и миграции', 'Пользователи, версии, состояния, evidence, аудит и usage'],
        ['Queue', 'Redis и BullMQ', 'Долгие и повторяемые операции с лимитами'],
        ['Files', 'S3 совместимый storage', 'Файлы попыток, экспорты и сертификаты'],
        ['Scanner', 'ClamAV compatible service', 'Карантин и антивирусная проверка'],
        ['Proxy', 'HTTPS reverse proxy', 'TLS, routing, security headers и публичные endpoints'],
    ], [1.25, 2.3, 3.5]
)

heading('Доменная модель и обязательные правила', 1)
body('Основные сущности должны быть версионными. Нельзя перезаписывать историю профиля, проблемы, курса или проверки.')
bullets([
    'User, ProfileVersion, Problem и ProblemVersion.',
    'Goal, DiagnosticSession, DiagnosticResult и DiagnosticContext.',
    'Course, CourseVersion, RouteVersion, Step и TaskVersion.',
    'Attempt, ReviewVersion, Evidence и CompetencyEstimate.',
    'BuildRequest, PublishDecision и DecisionEvent.',
    'SourceRecord, MethodVersion, Consent, Certificate и AuditEvent.',
    'Job, InboundMaxEvent, NotificationOutbox и AiUsageLedger.'
])
body('Состояния обучения, фоновой задачи, попытки и спора должны быть разными state machine. Техническое освоение требует двух независимых PASS по заданиям разных семейств без раскрытого решения, без помощи и без открытого спора. Evidence имеет срок действия и может быть отозван.')

heading('Что переносится из VibeWork', 1)
table(
    ['Источник', 'Решение', 'Способ использования'],
    [
        ['Профиль и нормализация', 'Перенести логику', 'Сократить поля и переписать на TypeScript'],
        ['Режимы анализа', 'Перенести идеи', 'Использовать для классификации четырёх типов запроса'],
        ['Каталог 32 ресурсов', 'Перенести данные после ревизии', 'Добавить source status provenance license checked at и catalog version'],
        ['Шесть путей обучения', 'Перенести как черновик', 'Преобразовать в LearningBlock и prerequisite graph'],
        ['LLM prompts и fallback', 'Перенести паттерны', 'Создать provider abstraction и schema validation'],
        ['HH и MTS', 'Оставить вторичным модулем', 'Показывать после подтверждённой цели'],
        ['Симулятор', 'Адаптировать', 'Использовать как практическое или readiness задание'],
        ['File validation', 'Переписать', 'Сохранить проверки MIME имени квот и owner'],
        ['Docker Nginx env docs', 'Использовать как ориентир', 'Создать новую многосервисную конфигурацию'],
        ['Responsive CSS', 'Перенести требования', 'Реализовать заново на React и MAX UI'],
    ], [1.55, 1.8, 3.7]
)
body('Отобранные материалы уже скопированы в /Users/inkoromi21/Documents/VibeWorkMax/legacy/vibework_reference. В README описаны ограничения, а SOURCE_MANIFEST.md связывает каждую группу файлов с задачами нового проекта. Эта выборка является первым источником для промтов; полный VibeWork нужно открывать только при доказанной нехватке материала.')

heading('Что не переносится из VibeWork', 1)
bullets([
    'FastAPI и существующие Python роуты.',
    'SQLite и текущая схема хранения.',
    'Email регистрация как основной способ входа.',
    'Монолитная страница и старая навигация.',
    'Обязательные карьерные тесты перед решением учебной проблемы.',
    'Ручной статус done как доказательство знания.',
    'Текущие API paths, не совпадающие с контрактами комплекта 4.0.',
    'Свободная генерация плана без каталога, версий и валидаторов.',
    'Текущая простая админка пользователей.',
    'Неотслеживаемые данные из website data и пользовательская SQLite база.'
])
body('Организационный риск зафиксирован в README выборки: в VibeWork есть незакоммиченные изменения, тестовый импорт GapBar сломан, один тест удаления аккаунта падает, а каталог website data игнорируется git. В выборку не попали секреты, пользовательская БД, runtime, кэши и старый монолитный интерфейс. Поэтому перенос выполняется только из явно перечисленных артефактов, а не развитием старого приложения на месте.')

heading('Использование Yandex AI Studio', 1)
body('Модель подходит для управляемой адаптации образовательного материала, но не должна самостоятельно определять освоение, выбирать неизвестные компетенции или публиковать курс. Решения остаются в коде, а модель работает только с разрешённым контекстом.')
figure(AI_PIPE, 'Рисунок 3  Контролируемый конвейер генерации курса', 'Конвейер генерации через Yandex AI с жёсткими фильтрами и валидацией', width=5.9)

heading('Рекомендуемые модели и стоимость', 2)
table(
    ['Модель', 'Роль', 'Вход за 1000', 'Выход за 1000', 'Контекст'],
    [
        ['Alice AI LLM Flash', 'Классификация, уточнение, варианты, объяснение', '0,10 ₽', '0,20 ₽', '65 536'],
        ['Alice AI LLM', 'Урок, сложная адаптация, неоднозначная проверка', '0,50 ₽', '1,20 ₽', '131 072'],
        ['YandexGPT Pro 5.1', 'Benchmark challenger', '0,80 ₽', '0,80 ₽', '32 768'],
        ['YandexGPT Lite 5', 'Резервный дешёвый вариант', '0,20 ₽', '0,20 ₽', '32 768'],
    ], [1.85, 2.35, 1.05, 1.05, 1.0]
)
body('Типичный гибридный сценарий оценивается примерно в 12,15 рубля, а с резервом 20 процентов на повторные запросы — в 14,58 рубля. Сто сценариев — около 1458 рублей, тысяча — около 14 580 рублей. Точные значения нужно записывать по usage каждого ответа.')
body('Начальные ограничители для демонстрационной среды: MAX_OPERATION_COST равен 25 рублям, DAILY_AI_BUDGET равен 1500 рублям. Их должен утвердить тимлид. Для запросов с пользовательскими данными нужен заголовок x data logging enabled false и режим store false, если используется Responses API.')

heading('Этапы реализации', 1)
figure(ROADMAP, 'Рисунок 4  Рекомендуемый порядок разработки', 'Дорожная карта из восьми последовательных фаз')
table(
    ['Этап', 'Результат', 'Оценка разработчика'],
    [
        ['Основа', 'Чистый репозиторий, контракты и решения', '1–2 дня'],
        ['Платформа', 'Монорепо, инфраструктура, БД и очереди', '4–5 дней'],
        ['MAX', 'Основной бот, mini app auth и webhook', '4–6 дней'],
        ['Диагностика', 'Проблема, вопросы, результат и цель', '4–6 дней'],
        ['Курс и AI', 'Каталог, build, Yandex и публикация', '5–7 дней'],
        ['Обучение', 'Попытки, review, evidence и replan', '5–7 дней'],
        ['Полнота', 'Файлы, оператор, сертификат и данные', '4–6 дней'],
        ['Сдача', 'E2E, нагрузка, безопасность и пакет', '4–5 дней'],
    ], [1.2, 4.4, 1.65]
)

heading('Библиотека промтов для Codex', 1)
body('Следующие промты образуют полный путь создания проекта. Они намеренно ограничены. Один промт должен завершаться проверяемым изменением, а не общей просьбой переделать систему. В каждом промте есть отдельный набор статических, автоматических, негативных, интеграционных и smoke проверок, а также обязательный цикл исправления найденных проблем. Промты, использующие VibeWork, дополнительно содержат правила переписывания на целевой стек: переносится поведение и набор тестовых случаев, но не Python runtime, FastAPI, SQLite и старые клиенты внешних сервисов. После каждого промта сохраняйте результат отдельным commit, если рабочая ветка стабильна.')

# Prompt phases
heading('Фаза ноль Подготовка и границы', 1)
prompt_block('00', 'Аудит исходных материалов', 'Создать проверяемую карту требований и не начинать архитектуру на предположениях.', [
    'Проиндексируй документы в /Users/inkoromi21/Downloads и подготовленную выборку legacy/vibework_reference только для чтения.',
    'Создай docs/REQUIREMENTS_MAP.md с требованиями конкурса, тимлида, дополнения и контрактов.',
    'Для каждого требования укажи источник, приоритет, компонент, проверку и возможный конфликт.',
    'Отдельно зафиксируй, что бот является основным решением, а mini app расширяет его.',
    'Создай docs/SOURCE_INVENTORY.md со списком контрактов, фикстур и переносимых частей из SOURCE_MANIFEST.md; не сканируй полный VibeWork без зафиксированной причины.'
], ['Все FR 01–16 представлены в карте.', 'Отражены четыре типа пользовательского запроса.', 'Конфликты имеют явное решение и владельца.', 'Исходные файлы не изменены.'])

prompt_block('01', 'Фиксация VibeWork как источника', 'Зафиксировать переносимые материалы без копирования секретов, БД и старого runtime.', [
    'Прочитай legacy/vibework_reference/README.md и SOURCE_MANIFEST.md; используй зафиксированные там commit и состояние рабочей копии.',
    'Создай docs/VIBEWORK_REUSE.md с таблицей take adapt leave.',
    'Перечисли каталоги, алгоритмы, тестовые идеи и UX паттерны, которые можно перенести.',
    'Зафиксируй проблемы тестов GapBar, owner_user_id и игнорируемого website/data.',
    'Проверь подготовленную выборку и создай manifest с контрольными суммами, не копируя .env и vibework.db; полный VibeWork открывай только если конкретного источника нет в выборке.'
], ['Нет секретов и пользовательской БД.', 'Каждый переносимый элемент имеет назначение в новом продукте.', 'Python код не объявлен частью целевого runtime.'])

prompt_block('02', 'Реестр отсутствующих данных', 'Создать единый управляемый список ключей, доменов, контента и решений, которых пока нет.', [
    'Создай docs/MISSING_INPUTS.md с полями owner status deadline fallback и impact.',
    'Добавь разделы MAX, Yandex Cloud, S3, ClamAV, домен и TLS, контент, эксперт, юридические тексты, оператор, аналитика и submission.',
    'Для каждого отсутствующего секрета укажи только имя переменной окружения и способ безопасной передачи.',
    'Для каждого внешнего сервиса опиши mock или feature flag, позволяющий продолжить локальную разработку.'
], ['Ни одного придуманного ключа или URL.', 'Все блокеры имеют безопасный локальный fallback.', 'Файл пригоден как checklist тимлида.'])

heading('Фаза один Каркас и инфраструктура', 1)
prompt_block('03', 'Создание монорепозитория', 'Создать минимальный TypeScript monorepo для web api worker и общих пакетов.', [
    'Создай apps/web, apps/api, apps/worker и packages/contracts, domain, content, ai, shared.',
    'Настрой единый package manager, TypeScript strict, lint, format и test scripts.',
    'Добавь root README с командами разработки и картой директорий.',
    'Не добавляй feature код и не подключай внешние API.'
], ['Все пакеты компилируются.', 'Есть одна команда проверки типов и одна команда тестов.', 'Нет дублирующихся конфигураций без необходимости.'])

prompt_block('04', 'Импорт контрактов комплекта 4 0', 'Сделать предоставленные JSON Schema и OpenAPI источником типов и runtime validation.', [
    'Скопируй schemas.json, openapi-core.json и fixtures.json в packages/contracts с указанием происхождения.',
    'Добавь проверку контрольных сумм исходных копий.',
    'Настрой Ajv для 21 определения и генерацию TypeScript типов.',
    'Добавь тест, который проверяет 17 путей OpenAPI и все примеры схем.',
    'Не исправляй контракты молча. Расхождения записывай в docs/CONTRACT_DEVIATIONS.md.'
], ['Контрактные тесты проходят.', 'Типы импортируются api и web.', 'Изменение исходной схемы обнаруживается тестом.'])

prompt_block('05', 'Единые ошибки и идентификаторы', 'Ввести общие типы идентификаторов, ошибок, trace id и idempotency key.', [
    'Реализуй branded identifiers для основных сущностей без раскрытия последовательных ID наружу.',
    'Реализуй ApiError из схемы и Fastify error handler.',
    'Добавь correlation id на входящий запрос, job и вызов провайдера.',
    'Добавь redaction чувствительных полей в логах.'
], ['Ошибки соответствуют контракту.', 'Trace id возвращается клиенту и связывает логи.', 'Токены, initData и ответы диагностики не попадают в логи.'])

prompt_block('06', 'Docker Compose окружение', 'Поднять воспроизводимое локальное окружение без внешних ключей.', [
    'Добавь сервисы proxy, web, api, worker, postgres, redis, minio и clamav.',
    'Добавь healthcheck и зависимости готовности.',
    'Создай .dockerignore и .env.example без значений секретов.',
    'Добавь profile для облегчённого запуска без ClamAV, но production profile должен его включать.',
    'Проверь время docker build и зафиксируй результат.'
], ['docker compose config валиден.', 'Локальный stack стартует без облачных ключей.', 'Данные хранятся в именованных volumes.', 'Секреты не встроены в image.'])

prompt_block('07', 'Миграции PostgreSQL', 'Создать основу схемы данных и повторяемые миграции.', [
    'Выбери один migration tool и зафиксируй решение в ADR.',
    'Создай таблицы users, sessions, profile_versions, problems, problem_versions, goals и consents.',
    'Добавь UUID, created_at, updated_at и version поля там, где они нужны.',
    'Добавь owner indexes и ограничения уникальности.',
    'Добавь команду migrate и integration test на чистой БД.'
], ['Миграции применяются с нуля.', 'Повторный запуск безопасен.', 'Откат или compensating strategy описаны.', 'SQLite не используется.'])

prompt_block('08', 'Очередь фоновых заданий', 'Создать надёжный BullMQ контур для долгих операций.', [
    'Определи JobState отдельно от доменных состояний.',
    'Реализуй producer в api и consumer в worker.',
    'Добавь idempotency key, attempts, backoff, timeout и dead letter policy.',
    'Добавь endpoint чтения статуса по контракту JobAccepted.',
    'Добавь тест повторной доставки одного job.'
], ['Повторная обработка не создаёт дубликаты.', 'API не ждёт долгий job синхронно.', 'Состояние ошибки содержит безопасную причину и возможность повтора.'])

prompt_block('09', 'Outbox и аудит', 'Гарантировать согласованность доменного изменения и будущего уведомления.', [
    'Создай decision_events, audit_events и notification_outbox.',
    'Записывай доменное событие и outbox в одной транзакции.',
    'Реализуй worker отправки с блокировкой строк и повтором.',
    'Добавь дедупликацию по event_id.',
    'Запрети хранение сырых ответов пользователя в продуктовой аналитике.'
], ['Сбой отправки не откатывает доменное изменение.', 'Повторный worker не отправляет уведомление дважды.', 'Аудит содержит actor action target version result.'])

heading('Фаза два MAX бот как основной продукт', 1)
prompt_block('10', 'Приём webhook MAX', 'Реализовать безопасный и быстрый endpoint входящих событий MAX.', [
    'Создай HTTPS webhook route и типы Update.',
    'Проверь X-Max-Bot-Api-Secret constant time сравнением.',
    'Сохрани исходный event id и минимально нужный payload до ответа.',
    'Верни 200 после commit, а обработку передай в BullMQ.',
    'Добавь дедупликацию и метрику задержки.'
], ['Неверный secret отклоняется.', 'Дубликат не создаёт второе действие.', 'Ответ укладывается в 30 секунд.', 'Сырые токены не логируются.'], ['MAX_BOT_TOKEN', 'MAX_WEBHOOK_SECRET', 'Публичный HTTPS URL webhook'])

prompt_block('11', 'Клиент MAX Bot API', 'Создать ограниченный клиент актуального MAX API для сообщений и callbacks.', [
    'Используй platform-api2.max.ru и Authorization header.',
    'Реализуй send message, edit message, answer callback и upload при необходимости.',
    'Добавь timeout, rate limit ниже 30 rps, backoff для 429 и 503 и circuit breaker.',
    'Нормализуй ошибки провайдера и сохраняй request id без чувствительных данных.',
    'Добавь mock transport для тестов.'
], ['Нет токена в query параметрах.', 'Повтор допускается только для безопасных операций.', 'Mock покрывает success timeout 429 503 и malformed response.'], ['MAX_BOT_TOKEN'])

prompt_block('12', 'Диалоговая state machine бота', 'Создать независимый полезный сценарий бота на общей доменной модели.', [
    'Определи BotConversationState отдельно от LearningState.',
    'Поддержи вход, проблему, уточнение, подтверждение, диагностику, результат, цель, обучение, попытку, ожидание и продолжение.',
    'Каждый callback должен содержать короткий opaque payload и проверку текущего состояния.',
    'Добавь команды начать, продолжить, прогресс, помощь и отмена.',
    'Не храни состояние только в памяти процесса.'
], ['Рестарт процесса не теряет диалог.', 'Устаревший callback безопасно отклоняется.', 'Любой экран имеет понятный следующий шаг и выход.'])

prompt_block('13', 'Онбординг и согласия в боте', 'Дать пользователю полноценное начало пути без обязательного открытия приложения.', [
    'Обработай bot_started и start.',
    'Покажи краткое назначение, возрастные ограничения и ссылки на документы.',
    'Запроси необходимые согласия и сохрани их версию.',
    'Предложи начать новую проблему или продолжить незавершённую.',
    'Добавь кнопку открытия mini app как дополнительный способ, а не единственный путь.'
], ['Без согласия запрещён переход к обработке персональных данных.', 'Версия текста согласия сохраняется.', 'Повторный пользователь видит актуальное состояние.'], ['Тексты согласий и политика обработки данных', 'Возрастные правила и необходимость согласия представителя'])

prompt_block('14', 'Ввод и подтверждение проблемы в боте', 'Получить свободный запрос и подтвердить, что система поняла его правильно.', [
    'Прими свободный текст с ограничением длины.',
    'Создай Problem и ProblemVersion.',
    'Классифицируй запрос детерминированным fallback или AI job.',
    'Покажи пользователю краткую интерпретацию и кнопки подтвердить, изменить, начать заново.',
    'Не запускай диагностику до подтверждения.'
], ['Исходный текст и интерпретация версионны.', 'Изменение создаёт новую версию.', 'Неопределённый тип приводит к уточнению, а не случайному выбору.'])

prompt_block('15', 'Диагностика внутри бота', 'Позволить пройти значимую адаптивную диагностику непосредственно в диалоге.', [
    'Отправляй один PublicQuestion на сообщение.',
    'Используй inline кнопки для single и multi choice, отдельные не знаю, пропустить и пауза.',
    'Поддержи short answer и текстовую practical задачу.',
    'Сохраняй AnswerSubmission идемпотентно.',
    'Показывай адаптивный прогресс без обещания фиксированного числа вопросов.',
    'После восьми вопросов предложи завершить или добровольно продолжить до двенадцати.'
], ['Диагностика проходится без mini app.', 'Skip не считается ошибкой.', 'Повторный callback не создаёт второй ответ.', 'Пауза и продолжение работают после рестарта.'])

prompt_block('16', 'Результат и цель внутри бота', 'Выдать понятный результат и получить подтверждение цели в основном канале.', [
    'Отправь краткое объяснение known gap unknown и первый рекомендуемый шаг.',
    'Покажи источник результата и ограничения точности без псевдопроцентов.',
    'Дай кнопки подтвердить цель, изменить цель, посмотреть подробнее в mini app.',
    'Сохрани Goal только после подтверждения.',
    'Запусти course build job и сообщи ожидаемое состояние.'
], ['Результат понятен без приложения.', 'Цель версионна.', 'Пользователь видит, что генерация не мгновенна.', 'Повторное подтверждение идемпотентно.'])

prompt_block('17', 'Короткий урок внутри бота', 'Сделать бот самостоятельным образовательным интерфейсом для простых шагов.', [
    'Покажи название шага, причину включения, цель и короткий учебный материал.',
    'Разбей длинный текст на несколько сообщений с управляемой пагинацией.',
    'Дай кнопки пример, следующий фрагмент, практика, открыть полный урок.',
    'Сохраняй step_opened и позицию пользователя.',
    'Не показывай решение практики до ответа.'
], ['Короткий шаг полностью читается в боте.', 'Позиция синхронизируется с mini app.', 'Нет дублирования analytics events.'])

prompt_block('18', 'Практика и feedback внутри бота', 'Позволить решить простое задание, получить проверку и повторить попытку.', [
    'Поддержи ответы кнопкой и текстом.',
    'Создай Attempt с idempotency key.',
    'Покажи статус обработки, затем краткий Review по критериям.',
    'Дай кнопки попробовать снова, объяснить ошибку, оспорить и открыть подробности.',
    'Если решение было раскрыто, пометь попытку непригодной для independent pass.'
], ['Попытка создаётся один раз.', 'Feedback не раскрывает полный ответ преждевременно.', 'Спор создаётся из бота.', 'Результат синхронен с mini app.'])

prompt_block('19', 'Возврат и уведомления бота', 'Возвращать пользователя в правильный шаг после фоновых событий и перерывов.', [
    'Обработай plan published, attempt graded, route changed, dispute resolved, certificate ready.',
    'Создай шаблоны коротких уведомлений без персональных данных.',
    'Добавь кнопку продолжить с opaque start parameter.',
    'Уважай notification preferences и quiet hours, если они заданы.',
    'Добавь dedupe по notification event id.'
], ['Одинаковое событие не отправляется дважды.', 'Ссылка открывает правильный объект.', 'При недоступном mini app остаётся полезное действие в боте.'])

prompt_block('20', 'Помощь и восстановление сценария бота', 'Сделать диалог устойчивым к неизвестным сообщениям и потерянному контексту.', [
    'Обработай неизвестный текст, устаревшую кнопку и недоступное действие.',
    'Добавь помощь с текущим состоянием, а не общий список команд.',
    'Дай безопасный способ отменить текущий ввод и вернуться к последнему сохранённому шагу.',
    'Добавь correlation id в сообщение о технической ошибке.',
    'Не показывай stack trace и данные провайдера.'
], ['Из любого состояния есть восстановление.', 'Ошибки не повреждают доменное состояние.', 'Есть тесты неизвестного события и устаревшего callback.'])

heading('Фаза три Мини приложение', 1)
prompt_block('21', 'React оболочка MAX', 'Создать native looking оболочку mini app с MAX Bridge.', [
    'Подключи React, Vite, MAX UI и MAX Bridge.',
    'Создай routing, app shell и error boundary.',
    'Определи capabilities по platform и version.',
    'Добавь безопасные fallback для неподдерживаемых bridge методов.',
    'Не используй initDataUnsafe для доверия к пользователю.'
], ['Приложение открывается внутри MAX и обычного браузера в dev режиме.', 'Mobile web desktop состояния различимы.', 'Нет обязательного bridge метода без fallback.'])

prompt_block('22', 'Серверная сессия mini app', 'Проверить initData и создать защищённую пользовательскую сессию.', [
    'Передай исходную строку initData серверу.',
    'Отклони повторяющиеся параметры и отсутствие hash.',
    'Проверь HMAC SHA256 по алгоритму MAX.',
    'Проверь auth_date с лимитом 60 минут и clock skew 60 секунд.',
    'Создай Secure HttpOnly SameSite cookie и CSRF защиту для state changing routes.',
    'Добавь dev auth только под явным feature flag.'
], ['Поддельные и старые данные отклоняются.', 'Bot token никогда не попадает в браузер.', 'Dev auth невозможно включить случайно в production.'])

prompt_block('23', 'Экран проблемы и минимального профиля', 'Собрать только данные, влияющие на следующий шаг.', [
    'Создай свободное поле проблемы и подтверждение интерпретации.',
    'Показывай профильные вопросы только по необходимости.',
    'Добавь autosave черновика и version conflict handling.',
    'Поддержи четыре типа запроса без обязательного выбора профессии.',
    'Синхронизируй состояние с действиями бота.'
], ['Путь знания не требует карьерного теста.', 'Черновик восстанавливается.', 'Конфликт версий не затирает новый ответ.'])

prompt_block('24', 'Интерфейс диагностики', 'Реализовать один вопрос на экран и полные состояния диагностики.', [
    'Поддержи single multi short practical preference.',
    'Добавь не знаю, пропустить, пауза, продолжить и завершить частично.',
    'Показывай адаптивный прогресс и причину дополнительного вопроса при необходимости.',
    'Раздели сохранение ответа и переход к следующему вопросу.',
    'Добавь keyboard и screen reader поведение.'
], ['Нет скрытого автоматического submit.', 'Все типы вопросов доступны с клавиатуры.', 'Обновление страницы продолжает с нужного вопроса.'])

prompt_block('25', 'Экран результата и цели', 'Показать объяснимый результат без ложной точности.', [
    'Отдельно отобрази known gap unknown и evidence sufficiency.',
    'Покажи, какие ответы повлияли на вывод, без раскрытия закрытых решений.',
    'Добавь первый шаг и причины рекомендации.',
    'Дай подтвердить или изменить цель.',
    'Добавь ссылку на методы и источники.'
], ['Нет смешения интереса и знания.', 'Нет неподтверждённых процентов mastery.', 'Изменение цели создаёт новую версию.'])

prompt_block('26', 'Маршрут и курс', 'Показать активную пару route и course с причинами каждого шага.', [
    'Создай обзор маршрута и экран шага.',
    'Покажи reason kind и текстовое объяснение.',
    'Раздели continue save и submit.',
    'Покажи версии и дату обновления без перегрузки интерфейса.',
    'Добавь состояние перестройки маршрута и возврат к старой активной версии до публикации новой.'
], ['Пользователь понимает, почему шаг включён.', 'Candidate не показывается как active.', 'Старая версия остаётся доступна до атомарного переключения.'])

prompt_block('27', 'Урок задание и review', 'Реализовать полный учебный шаг в mini app.', [
    'Покажи учебный материал, источники и практику.',
    'Скрой решение до отправки ответа.',
    'Поддержи текст и разрешённые файлы.',
    'Покажи rubric criteria, status и feedback.',
    'Добавь retry и dispute без изменения старой ReviewVersion.'
], ['Практика идёт до решения.', 'История попыток неизменяема.', 'Загрузка файла показывает карантин и результат проверки.'])

prompt_block('28', 'Прогресс сертификат и данные', 'Собрать финальные пользовательские экраны.', [
    'Покажи mastery state и evidence sufficiency по компетенциям.',
    'Отдельно покажи обычный прогресс прохождения и подтверждённое освоение.',
    'Добавь сертификаты, opportunities и новый запрос.',
    'Добавь настройки уведомлений, экспорт и удаление данных.',
    'Для скачивания в MAX используй Bridge downloadFile с browser fallback.'
], ['Done не равен mastery.', 'Удаление требует подтверждения и показывает последствия.', 'Сертификат доступен только после серверной проверки цели.'])

prompt_block('29', 'Адаптивность и доступность', 'Проверить mini app на целевых размерах и способах ввода.', [
    'Настрой layouts для 360, 390 и 1280 пикселей.',
    'Обеспечь targets не меньше 44 на 44.',
    'Проверь zoom текста, keyboard navigation, focus visible и reduced motion.',
    'Добавь loading error empty offline pause resume states.',
    'Запусти accessibility audit и исправь критические проблемы.'
], ['Нет горизонтального скролла на целевых ширинах.', 'Фокус не теряется после async обновления.', 'Ключевой сценарий проходит без мыши.'])

heading('Фаза четыре Диагностика и контент', 1)
prompt_block('30', 'Версионный каталог контента', 'Создать импортируемый каталог компетенций блоков заданий рубрик и источников.', [
    'Определи форматы Competency LearningBlock TaskTemplate Rubric SourceRecord и MethodVersion.',
    'Добавь catalog_version, status, source, license, checked_at и expert_approved.',
    'Реализуй import validate preview publish archive.',
    'Запрети использование DRAFT и ARCHIVED в реальном маршруте.',
    'Импортируй исходные fixtures как DEMO_SYNTHETIC без смены expert_approved.'
], ['Синтетика явно отмечена.', 'Невалидная ссылка или неизвестная компетенция блокирует publish.', 'Старые версии остаются воспроизводимыми.'])

prompt_block('31', 'Классификация проблемы', 'Определить один из четырёх типов запроса с объяснимым fallback.', [
    'Сначала реализуй детерминированные правила и состояние ambiguous.',
    'Подключи AI только для неоднозначного текста.',
    'Требуй structured output с type confidence reason clarification_needed.',
    'Не используй confidence как вероятность знания.',
    'Сохрани модель, prompt version и результат в DecisionEvent.'
], ['Все четыре типа покрыты тестами.', 'Неоднозначный запрос создаёт уточнение.', 'AI не может создать пятый неизвестный тип.'])

prompt_block('32', 'План диагностической сессии', 'Построить ограниченный план вопросов из подтверждённой проблемы.', [
    'Создай DiagnosticContext snapshot входных версий.',
    'Выбери стартовый approved template.',
    'Сформируй очередь областей и stopping rules на 5–8 минут.',
    'Установи main limit 8 и hard limit 12 с отдельным согласием.',
    'Не требуй профессиональный тест для knowledge gap.'
], ['Одинаковый snapshot создаёт воспроизводимый план.', 'План содержит reason для области.', 'Лимиты проверяются сервером.'])

prompt_block('33', 'Адаптивный выбор вопроса', 'Выбирать следующий вопрос по evidence без зацикливания на одной области.', [
    'Реализуй selector approved QuestionTemplateVersion.',
    'Не допускай больше двух подряд неудачных вопросов одной области.',
    'Учитывай skip unknown preference и fixed method.',
    'LLM может менять только формулировку и разрешённые параметры варианта.',
    'Добавь deterministic seed для тестов.'
], ['Selector не повторяет уже отвеченный instance.', 'Skip не понижает knowledge score.', 'Disclosed solution помечает evidence соответствующим образом.'])

prompt_block('34', 'Завершение диагностики', 'Построить объяснимый DiagnosticResult и DiagnosticContext.', [
    'Реализуй stopping conditions completed completed partial paused.',
    'Рассчитай known gap unknown и evidence sufficiency детерминированно.',
    'Создай объяснение на основе фактов, а не свободного домысла модели.',
    'Сохрани ссылки на question instances и answer submissions.',
    'Разреши reassess без изменения старого результата.'
], ['Результат воспроизводим.', 'Partial не маскируется как полный.', 'Reassess создаёт новую сессию и связь с предыдущей.'])

heading('Фаза пять Yandex AI Studio', 1)
prompt_block('35', 'Интерфейс AI provider', 'Отделить доменную логику от конкретного облачного API.', [
    'Определи AiRequest AiResult ModelPolicy Usage Cost и ProviderError.',
    'Поддержи generate structured и generate text.',
    'Добавь operation kind, prompt version, schema, deadline и token ceilings.',
    'Создай fake provider для unit tests.',
    'Не импортируй SDK Yandex в domain package.'
], ['Domain тестируется без сети.', 'Провайдер возвращает usage и request id.', 'Ошибки normalised в retryable и final.'])

prompt_block('36', 'Адаптер Yandex AI Studio', 'Подключить Alice AI через OpenAI совместимый API безопасным серверным способом.', [
    'Используй base URL https://ai.api.cloud.yandex.net/v1.',
    'Поддержи Authorization Api Key и folder id.',
    'Добавь URI aliceai llm flash и aliceai llm.',
    'Добавь x data logging enabled false и store false где применимо.',
    'Установи timeout не больше 30 секунд и не больше одного retry.',
    'Логируй только метаданные usage latency request id.'
], ['Ключ не попадает в клиент и логи.', 'Есть mock contract test request shape.', '429 503 timeout и invalid JSON обработаны.'], ['YANDEX_FOLDER_ID', 'YANDEX_API_KEY', 'Подтверждение разрешённых моделей в каталоге Yandex Cloud'])

prompt_block('37', 'Бюджет и учёт токенов', 'Останавливать дорогие операции до вызова и сверять стоимость после ответа.', [
    'Создай AiUsageLedger и таблицу актуальных тарифов по конфигурации.',
    'Реализуй предварительную оценку входа и максимального выхода.',
    'Проверь MAX_OPERATION_COST и DAILY_AI_BUDGET перед enqueue и перед вызовом.',
    'После ответа сохрани фактический usage и рассчитанную цену.',
    'При превышении используй approved fallback и создавай audit event.'
], ['Гонка двух jobs не превышает дневной бюджет.', 'Цена считается отдельно по входу и выходу.', 'Изменение тарифов не требует изменения domain кода.'], ['Утверждённые MAX_OPERATION_COST и DAILY_AI_BUDGET'])

prompt_block('38', 'Версионные промпты и structured output', 'Сделать каждый тип генерации ограниченным схемой и воспроизводимым.', [
    'Создай отдельные prompt templates для clarification question variant explanation lesson feedback и replan.',
    'Раздели trusted instructions approved content и untrusted user input.',
    'Запроси JSON Schema output там, где модель это поддерживает.',
    'Проверь output через Ajv и доменные валидаторы.',
    'Разреши один repair retry, затем fallback.'
], ['Prompt injection из ответа пользователя не меняет системные правила.', 'Неизвестные IDs и URLs отклоняются.', 'Prompt version сохраняется с результатом.'])

prompt_block('39', 'Benchmark моделей', 'Выбрать модель по качеству задержке и цене на собственных fixtures.', [
    'Создай eval набор из expected cases и четырёх типов запросов.',
    'Сравни Alice AI LLM Flash, Alice AI LLM и YandexGPT Pro 5.1 на одинаковых inputs.',
    'Измерь schema validity, domain validity, source fidelity, rubric agreement, latency и cost.',
    'Выполни несколько прогонов с фиксированными настройками.',
    'Сохрани отчёт и ModelPolicy по operation kind.'
], ['Выбор модели основан на метриках.', 'Результаты не содержат исходные персональные данные.', 'Есть пороги автоматического отклонения.'], ['Доступ к выбранным моделям и достаточный тестовый бюджет'])

heading('Фаза шесть Персональный курс', 1)
prompt_block('40', 'BuildRequest и граф предпосылок', 'Сформировать неизменяемый snapshot и проверить prerequisite graph.', [
    'Собери ссылки на profile problem goal diagnostic evidence catalog policy и method versions.',
    'Проверь существование и совместимость всех версий.',
    'Проверь граф на циклы и неизвестные компетенции.',
    'Реализуй детерминированную топологическую сортировку.',
    'Сохрани hash BuildRequest.'
], ['Одинаковый snapshot имеет одинаковый hash.', 'Цикл блокирует build с понятной ошибкой.', 'Активные данные не читаются повторно во время build.'])

prompt_block('41', 'Детерминированный выбор маршрута', 'Выбрать разрешённые блоки до обращения к модели.', [
    'Примени hard filters возраста инструментов времени источников статуса и prerequisite.',
    'Используй reason kinds GAP PREREQUISITE CHECK EXPLORATION NEW GOAL PRACTICE FINAL ASSESSMENT.',
    'Выполни лексикографический выбор и tie breaking.',
    'Создай RoutePlan и candidate CourseVersion.',
    'Не позволяй LLM добавлять новые блоки.'
], ['Выбор воспроизводим.', 'Каждый шаг имеет reason.', 'Запрещённый источник никогда не попадает в prompt.'])

prompt_block('42', 'Адаптация урока и fallback', 'Адаптировать утверждённый блок под пользователя без изменения учебной истины.', [
    'Передай модели только разрешённый block, source fragments и constraints.',
    'Запрети изменение competencies answer keys rubrics и source ids.',
    'Проверь структуру, длину, возраст, инструменты, provenance и опасные инструкции.',
    'При первом сбое выполни один retry, затем используй готовый approved variant.',
    'Сохрани generation metadata.'
], ['Курс доступен при недоступном AI.', 'Fallback явно воспроизводим.', 'Сгенерированный текст не считается новым источником истины.'])

prompt_block('43', 'Атомарная публикация курса', 'Не публиковать устаревший candidate и переключать route course вместе.', [
    'Сравни BuildRequest versions с текущими версиями через compare and swap.',
    'Если данные изменились, пометь candidate SUPERSEDED.',
    'В одной транзакции создай PublishDecision и переключи active route course pair.',
    'Запиши DecisionEvent и outbox plan published.',
    'Добавь конкурентный integration test двух builds.'
], ['Старый job не затирает новый курс.', 'Route и course никогда не расходятся.', 'Повторный publish идемпотентен.'])

heading('Фаза семь Обучение и подтверждение результата', 1)
prompt_block('44', 'Попытки и версии проверки', 'Создать неизменяемую историю попыток и проверок.', [
    'Создай AttemptState и ReviewVersion отдельно.',
    'Проверяй objective ответы детерминированно.',
    'Для free text используй rubric evaluator с ограниченным AI feedback.',
    'Сохраняй критерии pass fail и evidence candidate.',
    'Новая проверка создаёт новую ReviewVersion, а не меняет старую.'
], ['Objective ответ не зависит от AI.', 'Review ссылается на конкретную rubric version.', 'Помощь и раскрытие решения учитываются.'])

prompt_block('45', 'Evidence и mastery', 'Реализовать правила подтверждения компетенции без ложной уверенности.', [
    'Создай Evidence со статусами valid revoked expired disputed.',
    'Реализуй mastery states unknown needs foundation developing demonstrated.',
    'Реализуй evidence sufficiency insufficient limited sufficient.',
    'Для demonstrated потребуй два PASS разных family id с assistance none и без открытого спора.',
    'Учти 90 дней и integrative task policy.'
], ['Один PASS не даёт mastery.', 'Повтор того же family не считается независимым.', 'Отозванное и просроченное evidence не засчитывается.'])

prompt_block('46', 'Перестройка маршрута', 'Создать новую версию маршрута после значимого evidence или изменения цели.', [
    'Определи события, запускающие replan.',
    'Сформируй новый BuildRequest с новыми версиями.',
    'Сохрани причину изменения и сравнение старого и нового маршрута.',
    'До публикации продолжай показывать текущий active route.',
    'После publish отправь пользователю объяснение через бот и mini app.'
], ['Нет перестройки на каждое незначимое действие.', 'Старые версии доступны для аудита.', 'Пользователь видит причину изменения.'])

prompt_block('47', 'Споры и оператор', 'Дать пользователю путь оспаривания без автоматического затирания оценки.', [
    'Создай DisputeState и очередь оператора.',
    'Разреши спор только владельцу попытки в допустимом состоянии.',
    'Сохрани сообщение пользователя и snapshot review.',
    'Оператор создаёт новую ReviewVersion и решение с причиной.',
    'Если оператор недоступен, предложи независимое fallback задание.'
], ['Открытый спор блокирует mastery evidence.', 'Все действия оператора аудируются.', 'Старый review остаётся неизменным.'], ['Список операторов и правила SLA', 'Тексты уведомлений о споре'])

heading('Фаза восемь Файлы безопасность и управление', 1)
prompt_block('48', 'Интерфейс файлового хранилища', 'Хранить файлы через одинаковый local и S3 interface.', [
    'Определи FileStorage put get delete signed url и metadata.',
    'Реализуй local adapter для разработки и S3 adapter для production.',
    'Используй непрозрачные storage keys и owner authorization.',
    'Ограничь размеры количество и разрешённые типы.',
    'Не делай bucket публичным.'
], ['Файл другого пользователя недоступен.', 'Оригинальное имя не используется как путь.', 'Удаление учитывает metadata и bytes.'], ['S3_ENDPOINT', 'S3_ACCESS_KEY', 'S3_SECRET_KEY', 'S3_BUCKET', 'S3_REGION'])

prompt_block('49', 'Карантин и ClamAV', 'Не выдавать загруженный файл до проверки содержимого.', [
    'Создай статусы uploading quarantined scanning clean infected rejected.',
    'Проверь extension MIME magic bytes и decompression limits до сканера.',
    'Передай scanning в worker.',
    'Разреши download только clean файлам.',
    'Удаляй или изолируй infected и записывай audit event.'
], ['Есть тест EICAR.', 'Сбой сканера не помечает файл clean.', 'Повтор scanning идемпотентен.'], ['CLAMAV_HOST', 'CLAMAV_PORT'])

prompt_block('50', 'Сертификат PDF', 'Генерировать сертификат только после серверной проверки достижения цели.', [
    'Определи неизменяемые данные Certificate.',
    'Проверь completion policy и отсутствие открытых конфликтов.',
    'Сгенерируй PDF на сервере и сохрани в FileStorage.',
    'Добавь hash и verification identifier.',
    'Отправь certificate ready через outbox.'
], ['Клиент не может сам выдать сертификат.', 'Повтор генерации не создаёт разные сертификаты без новой версии.', 'Скачивание работает через MAX Bridge и браузер.'], ['Утверждённый шаблон сертификата', 'Название организации и подписанты'])

prompt_block('51', 'Экспорт удаление и retention', 'Реализовать жизненный цикл пользовательских данных.', [
    'Создай asynchronous export job с машиночитаемым архивом.',
    'Создай deletion workflow с подтверждением, отменяемым grace period если он утверждён, и audit trail.',
    'Учти файлы, sessions, notifications, analytics identifiers и provider resources.',
    'Добавь retention jobs и отчёт удалённых категорий.',
    'Опиши backup RPO 24 и restore RTO 4.'
], ['Owner может запросить экспорт и удаление.', 'Удаление не падает на таблицах с другим именем owner column.', 'Политика retention тестируется на датах.'], ['Утверждённые сроки хранения', 'Юридические исключения и grace period'])

prompt_block('52', 'Операторская панель и RBAC', 'Создать минимальную безопасную панель для споров контента jobs и аудита.', [
    'Определи роли operator content editor admin auditor.',
    'Реализуй deny by default и server side checks.',
    'Добавь очереди disputes failed jobs content review и source review.',
    'Добавь publish archive и rollback каталога.',
    'Аудируй чтение чувствительных данных и все изменения.'
], ['Обычный пользователь не видит admin routes.', 'Content editor не управляет операторами.', 'Критические действия требуют подтверждения.'], ['Список ролей и первоначальных операторов'])

heading('Фаза девять Перенос полезных частей VibeWork', 1)
prompt_block('52A', 'Карта переписывания VibeWork на целевой стек', 'До переноса кода создать исчерпывающую карту замены технологий и границ повторного использования.', [
    'Прочитай SOURCE_MANIFEST.md и перечисленные в нём файлы только из legacy/vibework_reference.',
    'Создай docs/VIBEWORK_PORTING_MATRIX.md с колонками source, переносимое поведение, что отбросить, target package, target technology, fixtures, проверки и статус.',
    'Для каждого Python модуля укажи конкретную TypeScript точку назначения; для SQLite, requests, старого LLM и локальных файлов укажи целевые adapters и repositories.',
    'Добавь автоматическую проверку, которая не допускает Python runtime зависимости, FastAPI и SQLite в apps и packages нового проекта, оставляя legacy папку read only.',
    'На этом шаге не переноси feature код: подготовь карту, тестовые fixtures и небольшие validation scripts, необходимые следующим промтам.'
], ['Все строки SOURCE_MANIFEST.md покрыты.', 'У каждого переносимого поведения есть целевой TypeScript модуль и тест.', 'Нет пункта с формулировкой перенести как есть.', 'Запрещённые технологии явно сопоставлены с заменами.'])

prompt_block('53', 'Перенос учебного каталога', 'Преобразовать ресурсы и пути VibeWork в новый версионный формат.', [
    'Прочитай learning_catalog.json и learning_paths.json только из legacy/vibework_reference/website/data.',
    'Создай TypeScript migration CLI в packages/content для нового версионного content format.',
    'Добавь source status license checked at language age range и expert approved.',
    'Не публикуй записи без проверки источника.',
    'Добавь отчёт dropped transformed pending review.'
], ['32 ресурса учтены в отчёте.', 'Ни один внешний URL не становится PUBLISHED автоматически.', 'Script повторяем и тестируется fixture.'], ['Результат экспертной и лицензионной проверки ссылок'])

prompt_block('54', 'Перенос ветки выбора направления', 'Использовать карьерную логику VibeWork только для соответствующего запроса.', [
    'Перепиши правила school vocational career modes как TypeScript domain policy в packages/domain.',
    'Отдели preference signals от knowledge evidence.',
    'Сократи профиль до вопросов, влияющих на следующий выбор.',
    'Покажи несколько направлений с объяснением, а не единственный вердикт.',
    'Добавь тест, что knowledge gap не запускает карьерный опрос.'
], ['Нет смешения предпочтений и освоения.', 'Ветка активируется только для direction choice.', 'Старые психологические названия не выдаются как диагноз.'])

prompt_block('55', 'Перенос opportunities', 'Добавить HH и MTS после подтверждённой цели, не в основной учебный путь.', [
    'Перепиши фильтры и role mappings в отдельный TypeScript Opportunities module с provider interfaces.',
    'Раздели реальные интеграции и mock data.',
    'Маркируй источник, время обновления и ограничения.',
    'Добавь feature flags и graceful fallback.',
    'Не блокируй сертификат и обучение недоступностью opportunities.'
], ['Mock явно маркирован.', 'Сбой внешнего API не влияет на core.', 'Результаты фильтруются по возрастным и юридическим ограничениям.'], ['HH_USER_AGENT и подтверждение правил использования HH', 'Решение о реальной интеграции MTS или статическом каталоге'])

prompt_block('56', 'Перенос симулятора', 'Адаптировать workday simulator как практику или проверку готовности.', [
    'Выбери конкретную competency и rubric для каждого сценария.',
    'Преобразуй шаги симулятора в TypeScript versioned tasks и сохрани состояние через PostgreSQL repository.',
    'Отдели тренировочный режим от independent assessment.',
    'Не выдавай training result за mastery.',
    'Добавь bot friendly текстовый вариант и расширенный mini app вариант.'
], ['Тренировка с подсказкой не создаёт independent evidence.', 'Сценарий воспроизводим.', 'Есть fallback без LLM.'], ['Экспертное утверждение сценариев и рубрик'])

prompt_block('56A', 'Перенос проверок файлов без Python runtime', 'Сохранить полезные негативные сценарии VibeWork, полностью переписав файловый контур под целевой стек.', [
    'Извлеки из profile_files.py только правила, ограничения и тестовые примеры; составь таблицу соответствия старого случая новому TypeScript тесту.',
    'Дополни TypeScript contract suite FileStorage случаями размера, расширения, MIME, magic bytes, owner authorization, path traversal и повреждённого файла.',
    'Проверь новую цепочку Fastify streaming upload, quarantine, S3 или local storage, BullMQ и ClamAV без вызова Python.',
    'Добавь интеграционный тест EICAR и отказа сканера; скачивание разрешай только после server side статуса clean.',
    'Зафиксируй отброшенные детали старой реализации и причины в docs/VIBEWORK_PORTING_MATRIX.md.'
], ['Каждое полезное старое правило представлено TypeScript тестом либо явно отклонено.', 'API и worker запускаются без Python.', 'В новом коде нет SQLite, Pillow, UploadFile и Python subprocess.', 'Local и S3 adapters проходят одинаковый contract suite.'])

heading('Фаза десять Проверка и сдача', 1)
prompt_block('57', 'Контрактные тесты API', 'Проверить реализацию всех 17 путей и ошибок комплекта 4 0.', [
    'Сгенерируй запросы success и error для каждого operation.',
    'Проверь request response через JSON Schema.',
    'Проверь 202 JobAccepted и owner authorization.',
    'Проверь отсутствие старых assessment routes.',
    'Сформируй machine readable report.'
], ['Все 17 paths покрыты.', '400 401 403 404 409 и 422 имеют ApiError.', 'Контрактное расхождение не скрывается snapshot обновлением.'])

prompt_block('58', 'Unit тесты доменных правил', 'Зафиксировать критические инварианты диагностики курса и mastery.', [
    'Портируй expected cases комплекта 4 0.',
    'Добавь one pass not mastery, two independent pass mastery, same family rejected, disclosed solution rejected, revoked and expired evidence rejected.',
    'Добавь stale build superseded и atomic publish.',
    'Добавь skip not fail и no more than two failed questions per area.',
    'Используй deterministic clock и seed.'
], ['Каждый инвариант имеет отдельное понятное имя теста.', 'Тесты не используют сеть и реальное время.', 'Изменение правила требует явного обновления теста.'])

prompt_block('59', 'Интеграционные тесты отказов', 'Проверить устойчивость очередей хранилища MAX и Yandex.', [
    'Подними test containers Postgres Redis S3 и fake providers.',
    'Проверь timeout 429 503 invalid JSON duplicate webhook duplicate callback и worker crash.',
    'Проверь quarantine scanner failure и partial S3 failure.',
    'Проверь recovery после рестарта api и worker.',
    'Убедись, что retries не создают дубликаты.'
], ['Данные остаются согласованными.', 'Пользователь получает recoverable status.', 'Все повторяемые операции имеют idempotency tests.'])

prompt_block('60', 'E2E сценарий бота', 'Доказать, что основной путь работает в боте без mini app.', [
    'Создай test harness для входящих Updates и исходящих сообщений.',
    'Пройди start, consent, problem, confirm, diagnosis, result, goal, lesson, attempt, review и continue.',
    'Проверь pause resume и duplicate events.',
    'Проверь переход в mini app только как расширение.',
    'Сохрани краткий отчёт и transcript без персональных данных.'
], ['Главный сценарий завершается в боте.', 'Нет шага, где единственное действие равно открыть приложение.', 'Сообщения укладываются в ограничения MAX.'])

prompt_block('61', 'E2E mini app и синхронизация каналов', 'Проверить единое состояние между ботом и приложением.', [
    'Создай Playwright сценарии 360, 390 и 1280.',
    'Начни действие в боте и продолжи в mini app.',
    'Отправь попытку в mini app и получи уведомление в боте.',
    'Проверь browser fallback для Bridge.',
    'Проверь reload offline retry и session expiry.'
], ['Состояния совпадают в обоих каналах.', 'Нет двойных analytics events.', 'Ключевой путь проходит на mobile и web.'])

prompt_block('62', 'Нагрузочная проверка', 'Проверить установленные в ТЗ нагрузки и временные бюджеты.', [
    'Подготовь сценарий 20 активных сессий и 5 rps обычного API на 15 минут.',
    'Подготовь 5 одновременных course generation jobs с fake и staging provider.',
    'Измерь p50 p95 p99 error rate queue wait и provider latency.',
    'Проверь p95 обычного API не больше 2 секунд и первый шаг не больше 60 секунд.',
    'Сохрани воспроизводимый отчёт.'
], ['Пороговые значения сравниваются автоматически.', 'Тест не расходует облачный бюджет без явного флага.', 'Bottlenecks перечислены с доказательствами.'], ['Разрешение и бюджет для staging load test'])

prompt_block('63', 'Security review', 'Проверить границы доверия MAX пользователя контента файлов и LLM.', [
    'Проверь initData duplicate parameters old auth date HMAC timing webhook secret CSRF cookies и owner checks.',
    'Проверь prompt injection через problem answers sources filenames и operator text.',
    'Проверь SSRF path traversal unsafe download и MIME confusion.',
    'Проверь secret scanning dependencies и логирование.',
    'Сформируй findings с severity evidence и remediation.'
], ['Нет секретов в git и image.', 'Высокие finding исправлены или явно блокируют release.', 'Security tests воспроизводимы.'])

prompt_block('64', 'Документация и пакет запуска', 'Подготовить проект к независимому запуску жюри и тимлидом.', [
    'Обнови README архитектурой, быстрым стартом и demo path.',
    'Опиши зависимости, переменные, migrations, seed, mock и real modes.',
    'Подготовь Dockerfile Compose dockerignore env example и health checks.',
    'Добавь OpenAPI 3 1, DATA API yaml, тестовые данные и команды проверки.',
    'Проверь чистый clone и build не дольше пяти минут в целевой среде.'
], ['Новый человек запускает проект по README.', 'Все mock и synthetic данные маркированы.', 'Нет ссылок на локальную пользовательскую БД или секреты.'])

prompt_block('65', 'Конкурсные доказательства и демонстрация', 'Собрать технические материалы для доказательства ценности и работоспособности.', [
    'Добавь события problem confirmed diagnostic finished plan published step opened attempt graded route changed goal verified.',
    'Добавь dedupe event id и исключи raw answers.',
    'Подготовь три context passport и проверку второго каталога.',
    'Собери latency cost reliability и completion metrics.',
    'Подготовь скриншоты bot first сценария и данные для PDF презентации.'
], ['Метрики получены из системы, а не придуманы.', 'Пять интервью и пять наблюдений отмечены как внешняя зависимость, если не предоставлены.', 'Демо показывает основную ценность бота.'], ['Результаты пяти интервью', 'Результаты пяти наблюдаемых прогонов', 'Экспертное заключение', 'Финальный шаблон презентации и дедлайн'])

heading('Реестр внешних данных и блокеров', 1)
table(
    ['Категория', 'Что нужно получить', 'Что делать до получения'],
    [
        ['MAX', 'Bot token, bot name, webhook secret, verified profile, mini app URL', 'Mock transport, env placeholders, local signed fixtures'],
        ['Yandex', 'Folder ID, API key, model access, approved daily budget', 'Fake provider и recorded fixtures'],
        ['Домен', 'DNS, HTTPS certificate, public callback URLs', 'Local reverse proxy и documented placeholder'],
        ['S3', 'Endpoint, bucket, region, access key, secret key', 'MinIO adapter'],
        ['ClamAV', 'Production host, port and operational policy', 'Compose service и EICAR test'],
        ['Контент', 'Approved competencies, rubrics, answers, sources and licenses', 'DEMO SYNTHETIC fixtures with visible label'],
        ['Право', 'Consent text, privacy policy, age rules, retention periods', 'Versioned placeholders blocked from production'],
        ['Оператор', 'Accounts, roles, SLA and escalation rules', 'Seeded local operator under dev flag'],
        ['Opportunities', 'HH usage conditions and MTS integration decision', 'Static labelled demo catalog'],
        ['Бренд', 'Name, logo, certificate template and signer data', 'Neutral text only'],
        ['Конкурс', 'Deadline, submission form, presentation template', 'Maintain release checklist'],
        ['Исследование', 'Five interviews, five observed runs, expert review', 'Instrumentation and evidence templates'],
    ], [1.35, 3.2, 2.65]
)

heading('Критерии готовности конкурсного сценария', 1)
bullets([
    'Пользователь начинает и завершает основной сценарий в боте.',
    'Mini app подключён к боту и добавляет сложные визуальные возможности.',
    'Работа проверена в мобильном и веб клиенте MAX.',
    'Проблема подтверждена пользователем до диагностики.',
    'Диагностика адаптивна, ограничена по времени и различает skip unknown и fail.',
    'Результат объясним и не выдаёт интерес за знание.',
    'Курс построен из разрешённого версионного каталога.',
    'Практика предшествует показу решения.',
    'Оценка связана с конкретной рубрикой.',
    'Evidence и mastery вычисляются на сервере.',
    'Сбой Yandex не разрушает сценарий.',
    'Синтетика и mock интеграции явно маркированы.',
    'Секретов и персональных данных нет в репозитории и логах.',
    'Docker сборка воспроизводима.',
    'Есть OpenAPI, README, тестовые данные и презентация.'
])

heading('Контрольные решения перед началом production', 1)
numbered([
    'Тимлид подтверждает обязательность Node TypeScript и полный объём первой стадии.',
    'Команда подтверждает bot first пользовательский путь и объём действий без mini app.',
    'Предметный эксперт утверждает хотя бы один полный учебный каталог.',
    'Определяются тексты согласий и правила работы с несовершеннолетними.',
    'Утверждаются модели Yandex, MAX_OPERATION_COST и DAILY_AI_BUDGET.',
    'Выдаётся verified MAX bot и публичный HTTPS домен.',
    'Выбирается production S3 и политика антивирусного карантина.',
    'Назначаются операторы и контентные редакторы.',
    'Фиксируется freeze date и commit для конкурсной сдачи.'
])

heading('Источники и ссылки', 1)
source_line('ТЗ конкурса', '/Users/inkoromi21/Downloads/Obrazovatelnye_reshenia_1.pdf')
source_line('Дополнение по критериям', '/Users/inkoromi21/Downloads/Dopolnenie_k_TZ_MAX_po_kriteriam_khakatona.docx')
source_line('Комплект тимлида', '/Users/inkoromi21/Downloads/TZ_MAX_komplekt_4_0/')
source_line('Исходный проект', '/Users/inkoromi21/Documents/VibeWork/')
source_line('Документация mini apps MAX', 'dev.max.ru/docs/webapps/introduction', 'https://dev.max.ru/docs/webapps/introduction')
source_line('MAX Bridge', 'dev.max.ru/docs/webapps/bridge', 'https://dev.max.ru/docs/webapps/bridge')
source_line('Валидация MAX initData', 'dev.max.ru/docs/webapps/validation', 'https://dev.max.ru/docs/webapps/validation')
source_line('Webhook MAX', 'dev.max.ru/docs-api/methods/POST/subscriptions', 'https://dev.max.ru/docs-api/methods/POST/subscriptions')
source_line('Модели Yandex AI Studio', 'aistudio.yandex.ru/ru/docs/ai-studio/concepts/generation/models', 'https://aistudio.yandex.ru/ru/docs/ai-studio/concepts/generation/models')
source_line('Тарифы Yandex AI Studio', 'aistudio.yandex.ru/ru/docs/ai-studio/pricing', 'https://aistudio.yandex.ru/ru/docs/ai-studio/pricing')
source_line('Отключение логирования Yandex', 'aistudio.yandex.ru/ru/docs/ai-studio/operations/disable-logging', 'https://aistudio.yandex.ru/ru/docs/ai-studio/operations/disable-logging')

# Insert static clickable TOC entries before marker
parent = toc_marker._p.getparent()
idx = parent.index(toc_marker._p)
for level, title, anchor in toc_entries:
    p_el = OxmlElement('w:p')
    p_pr = OxmlElement('w:pPr')
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:after'), '38')
    p_pr.append(spacing)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), str((level-1)*360))
    p_pr.append(ind)
    p_el.append(p_pr)
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('w:anchor'), anchor)
    rr = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1F5A8A')
    rpr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20' if level == 1 else '18')
    rpr.append(sz)
    if level == 1:
        b = OxmlElement('w:b')
        rpr.append(b)
    rr.append(rpr)
    tt = OxmlElement('w:t')
    tt.text = title
    rr.append(tt)
    hl.append(rr)
    p_el.append(hl)
    parent.insert(idx, p_el)
    idx += 1
parent.remove(toc_marker._p)

# Core properties
doc.core_properties.title = 'Образовательная платформа MAX Книга разработчика'
doc.core_properties.subject = 'Архитектура bot first, перенос VibeWork и библиотека промтов Codex'
doc.core_properties.author = 'Команда проекта MAX'
doc.core_properties.keywords = 'MAX, бот, мини приложение, React, Fastify, Yandex AI Studio, Codex'

# Keep tables and set document defaults
settings = doc.settings._element
update_fields = OxmlElement('w:updateFields')
update_fields.set(qn('w:val'), 'true')
settings.append(update_fields)

doc.save(OUT)
print(OUT)
